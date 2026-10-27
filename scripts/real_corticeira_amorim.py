"""Corticeira Amorim, S.G.P.S., S.A. — consolidated statements FY2024/FY2023 copied as printed in the 2024
annual report (English, pages 377–380): quarterly columns before the label, dates as column headers and the
unit note at the foot of each statement, as an analyst pasting from the PDF would get them."""
import sys
from pathlib import Path
from openpyxl import Workbook

UNIT = "Values in thousand euros."
BS = [  # label, note, 31 Dec 2024, 31 Dec 2023
    ["Assets"],
    ["Tangible assets", 8, 435511, 438297], ["Intangible assets", 9, 15073, 18018], ["Right of use", 11, 5242, 5046],
    ["Goodwill", 9, 29165, 23872], ["Biological assets", 10, 4324, 4952], ["Investment property", 12, 2204, 2241],
    ["Investments in associates and joint ventures", 13, 35322, 32630], ["Other financial assets", 14, 1640, 2097],
    ["Deferred tax assets", 15, 20379, 20203], ["Other debtors", 18, 1518, 1895],
    ["Non-current assets", None, 550376, 549251],
    ["Inventories", 16, 466545, 516497], ["Biological assets", 10, 711, 1391], ["Trade receivables", 17, 194403, 203080],
    ["Income tax assets", 15, 19630, 7951], ["Other debtors", 14, 40558, 41726], ["Other current assets", 18, 13335, 21937],
    ["Cash and cash equivalents", 19, 76636, 73394],
    ["Current assets", None, 811818, 865974],
    ["Total Assets", None, 1362194, 1415225],
    ["Equity"],
    ["Share capital", 20, 133000, 133000], ["Other reserves", 20, 541588, 488311], ["Net Income", None, 69699, 88897],
    ["Non-Controlling Interest", 21, 90770, 89835], ["Total Equity", None, 835057, 800044],
    ["Liabilities"],
    ["Interest-bearing loans", 22, 119053, 101793], ["Other financial liabilities", 24, 6651, 8300],
    ["Other liabilities", 24, 2400, 0], ["Provisions", 25, 5691, 7942], ["Post-employment benefits", 26, 3210, 3228],
    ["Deferred tax liabilities", 15, 40586, 42715],
    ["Non-current liabilities", None, 177592, 163979],
    ["Interest-bearing loans", 22, 153270, 212440], ["Trade payables", 23, 112159, 159000],
    ["Other financial liabilities", 24, 51070, 51497], ["Other liabilities", 24, 28033, 24320],
    ["Income tax liabilities", 15, 5012, 3946],
    ["Current liabilities", None, 349545, 451203],
    ["Total Liabilities and Equity", None, 1362194, 1415225],
]
IS = [  # 4Q24, 4Q23, label, note, 2024, 2023
    [212813, 222248, "Sales", 7, 939061, 985467],
    [-103784, -95590, "Costs of goods sold and materials consumed", None, -443873, -485063],
    [4043, -11584, "Change in manufactured inventories", None, 1074, 7229],
    [-39702, -38532, "Third party supplies and services", 27, -151501, -151167],
    [-44432, -44616, "Staff costs", 28, -193191, -189732],
    [-1046, 476, "Impairments of assets", 29, -789, -1186],
    [6220, 5794, "Other income and gains", 30, 17849, 20088],
    [-4195, -995, "Other costs and losses", 30, -11064, -8651],
    [29917, 37200, "Operating profit before depreciation", None, 157566, 176985],
    [-15204, -13856, "Depreciation", "8, 9, 10, 11, 12", -57636, -51915],
    [14713, 23345, "Operating profit", None, 99930, 125070],
    [6730, 782, "Non-recurrent results", 31, 1434, 782],
    [-3057, -3439, "Financial costs", 32, -12865, -9319],
    [236, 511, "Financial income", 32, 973, 1494],
    [1164, -956, "Share of (loss)/profit of associates and joint ventures", 13, 4305, 3011],
    [19785, 20243, "Profit before tax", None, 93776, 121039],
    [2170, 4483, "Income tax", 15, -16630, -20903],
    [21955, 24725, "Profit after tax", None, 77146, 100136],
    [90, 2839, "Non-controlling Interest", 21, 7447, 11239],
    [21865, 21886, "Net Income attributable to the equity holders of Corticeira Amorim", None, 69699, 88897],
]
CF = [
    [None, None, "OPERATING ACTIVITIES"],
    [237377, 252125, "Collections from customers", None, 1004992, 1046035],
    [-160512, -208091, "Payments to suppliers", None, -725733, -898194],
    [-51942, -44205, "Payments to employees", None, -189968, -177020],
    [24923, -171, "Operational cash flow", None, 89291, -29179],
    [-16093, -9159, "Payments/collections - income tax", None, -30115, -27138],
    [17622, 18475, "Other collections/payments related with operational activities", None, 61737, 82614],
    [26452, 9145, "CASH FLOW FROM OPERATING ACTIVITIES", None, 120913, 26296],
    [None, None, "INVESTMENT ACTIVITIES"],
    [None, None, "Collections due to:"],
    [641, 2573, "Tangible assets", None, 1179, 4540],
    [41, 88, "Intangible assets", None, 87, 88],
    [18897, 0, "Disposal of Subsidiaries (net of cash and cash equivalents)", 6, 18915, 0],
    [0, 0, "Other financial assets", None, 0, 1841],
    [0, 7000, "Investment properties", None, 0, 7000],
    [433, 899, "Interests and similar gains", None, 1310, 2080],
    [350, 250, "Dividends", None, 1499, 1842],
    [None, None, "Payments due to:"],
    [-10224, -29952, "Tangible assets", None, -40397, -92651],
    [-4380, -5981, "Acquisition of Subsidiaries (net cash and cash equivalents)", 6, -4380, -6081],
    [228, -400, "Intangible assets", None, -1325, -2500],
    [5985, -25523, "CASH FLOW FROM INVESTMENTS ACTIVITIES", None, -23113, -83841],
    [None, None, "FINANCIAL ACTIVITIES"],
    [None, None, "Collections due to:"],
    [57200, 21238, "Loans", None, 109650, 129241],
    [2100, 1026, "Government grants", None, 7006, 1943],
    [0, -185, "Transactions with non-controlling interest", None, 0, 0],
    [531, 288, "Others", None, 1501, 2642],
    [None, None, "Payments due to:"],
    [-95631, -4950, "Loans", None, -114309, -66200],
    [-2617, -3533, "Interests and similar expenses", None, -12411, -8445],
    [-549, -805, "Leasing", None, -978, -1723],
    [-11970, -11970, "Dividends paid to company's shareholders", 20, -38570, -38570],
    [-1138, -2328, "Dividends paid to non-controlling interest", 21, -5723, -5493],
    [149, -1031, "Government grants", None, -2480, -3220],
    [-190, -210, "Others", None, -722, -823],
    [-52116, -2460, "CASH FLOW FROM  FINANCING ACTIVITIES", None, -57037, 9353],
    [-19679, -18838, "Change in cash and cash equivalents", None, 40763, -48192],
    [192, 159, "Exchange rate effect", None, 70, -17],
    [47451, 5810, "Cash and cash equivalents at beginning of period", 19, -12869, 35341],
    [27964, -12869, "Cash and cash equivalents at end of period", 19, 27964, -12869],
]

wb = Workbook()
src = wb.active
src.title = "Source"
src.append(['Corticeira Amorim, S.G.P.S., S.A. — published consolidated accounts FY2024'])
src.append(['Copied as printed from the 2024 annual report (English), pages 377–380. Figures in thousand euros.'])
src.append(["Source: " + 'https://www.amorim.com/xms/files/Investidores/Relatorio_Anual/Amorim_RC24_EN_low.pdf'])
src.append(["Real, published figures. Used to test the agent on real accounts."])
src.column_dimensions["A"].width = 110
ws = wb.create_sheet("Financial position")
ws.append(["Consolidated Statement of Financial Position"])
ws.append([])
ws.append([None, "Notes", "31 December 2024", "31 December 2023"])
for r in BS:
    ws.append(r)
ws.append([])
ws.append([UNIT])
for title, heading, rows in [("Income statement", "Consolidated income statement by nature", IS),
                             ("Cash flow", "Consolidated statement of cash flow", CF)]:
    ws = wb.create_sheet(title)
    ws.append([heading])
    ws.append([])
    ws.append(["4Q24", "4Q23", None, "Notes", 2024, 2023])
    ws.append(["(non audited)", "(non audited)"])
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append([UNIT])
out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / 'data/Corticeira_Amorim_2024_published.xlsx')
wb.save(out)
print("saved", out)
