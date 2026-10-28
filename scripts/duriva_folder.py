"""Duriva's data room as a seller would send it: five documents instead of one workbook.

    python scripts/duriva_folder.py data/Duriva_data_room.xlsx data/Duriva_data_room

Writes four Excel files (statements, sales by customer, management's adjusted EBITDA, debt schedule) with
every cell where it was in the single workbook, and the management report as a PDF. tests/test_dataroom.py
requires that reading the folder gives exactly the same figures and findings as reading the single file.
"""
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

FILES = {
    "01 Demonstrações financeiras 2023-2025.xlsx": ["Leia-me", "DRE", "Balanço", "DFC"],
    "02 Vendas por cliente.xlsx": ["Vendas por Cliente"],
    "03 EBITDA ajustado (management presentation).xlsx": ["EBITDA Ajustado"],
    "04 Mapa da dívida.xlsx": ["Dívida"],
}
REPORT, PDF_NAME = "Relatório de Gestão", "05 Relatório de gestão 2025.pdf"


def _escape(line: str) -> bytes:
    return line.encode("cp1252").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def _wrap(text: str, width: int = 92) -> list[str]:
    lines, line = [], ""
    for word in text.split():
        if len(line) + len(word) + 1 > width:
            lines.append(line)
            line = word
        else:
            line = f"{line} {word}".strip()
    return lines + [line] if line else lines


def write_pdf(path: Path, paragraphs: list[tuple[str, bool]]) -> None:
    """A plain text PDF (Helvetica, A4), with no dependency: enough for any PDF reader to extract the text."""
    rows = []
    for text, bold in paragraphs:
        rows += [(line, bold) for line in _wrap(text)] + [("", False)]
    pages = [rows[i:i + 50] for i in range(0, len(rows), 50)]
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", None,
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>"]
    kids = []
    for page in pages:
        stream = b"BT 11 TL 56 790 Td " + b" ".join(
            (b"/F2" if bold else b"/F1") + b" 10 Tf (" + _escape(line) + b") Tj T*" for line, bold in page) + b" ET"
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        objects.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents %d 0 R "
                       b"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> >>" % len(objects))
        kids.append(len(objects))
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % k for k in kids), len(kids))
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))


def main(source: str, folder: str) -> None:
    src = load_workbook(source, data_only=True)
    out_dir = Path(folder)
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, sheets in FILES.items():
        wb = Workbook()
        wb.remove(wb.active)
        for title in sheets:
            ws, new = src[title], wb.create_sheet(title)
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        new.cell(row=c.row, column=c.column, value=c.value)
            for col, dim in ws.column_dimensions.items():
                new.column_dimensions[col].width = dim.width
        wb.save(out_dir / name)
    report = src[REPORT]
    paragraphs = []
    for row in report.iter_rows(values_only=True):
        cells = [str(v) for v in row if v not in (None, "")]
        if len(cells) == 1:
            paragraphs.append((cells[0], True))
        elif cells:
            paragraphs += [(cells[0], True), (" ".join(cells[1:]), False)]
    write_pdf(out_dir / PDF_NAME, paragraphs)
    print("saved", ", ".join(sorted(p.name for p in out_dir.iterdir())))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
