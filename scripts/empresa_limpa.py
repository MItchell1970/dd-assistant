"""Cria a segunda empresa fictícia, "limpa": Velmora Material Elétrico, S.A.

Serve para medir alarmes falsos: uma empresa saudável, com demonstrações que batem entre
si, não deveria disparar regras. Gera dois ficheiros em data/:
    Velmora_data_room.xlsx    o data room, com as mesmas folhas do da Duriva
    DD_template_velmora.xlsx  o template com a folha Input preenchida e recalculada pelo Excel

Correr (Windows, com o Excel instalado, na pasta do projeto):
    python scripts/empresa_limpa.py

Limitação: a empresa foi desenhada por quem conhece as regras. Zero alertas aqui é um teste
de sanidade, não uma medida independente de alarmes falsos.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
TEMPLATE = DATA / "DD_template_modulo1.xlsx"
OUT_ROOM = DATA / "Velmora_data_room.xlsx"
OUT_TEMPLATE = DATA / "DD_template_velmora.xlsx"
NAME = "Velmora Material Elétrico, S.A."

PL = [2023, 2024, 2025]          # anos com DRE e DFC
BS = [2022, 2023, 2024, 2025]    # anos com balanço

# ---------------------------------------------------------------- dados base (€ milhares)
# Demonstração de resultados (gastos com sinal negativo)
vendas = [15200, 15900, 16700]
cmv = [-10944, -11385, -12041]
fse = [-1520, -1580, -1650]
pess = [-1410, -1470, -1540]
imp = [-30, -35, -35]
orend = [60, 55, 65]             # 2025 inclui 5 de ganho na venda de duas viaturas
ogast = [-40, -45, -50]
da = [-290, -300, -315]
juros = [-105, -90, -70]
irc = [-207, -236, -239]
colaboradores = [46, 47, 49]

# Balanço: rubricas de exploração e dívida (a caixa e o capital próprio saem dos fluxos)
aft_2022 = 2900
inv = [2170, 2250, 2340, 2470]
cli = [2480, 2582, 2700, 2840]
oac = [240, 250, 255, 265]
forn = [1750, 1800, 1870, 1980]
opc = [480, 500, 515, 535]
emprestimo = [2100, 1800, 1500, 1200]   # amortização de 300 por ano (dívida corrente: 300)
locacoes = [300, 240, 180, 120]         # amortização de 60 por ano (dívida corrente: 60)
caixa_2022, capital_social, rl_2022 = 700, 500, 650

# Investimento, dividendos e venda de ativos
capex = [-330, -350, -380]
alienacao = [0, 0, 20]
ganho_alienacao = [0, 0, 5]
dividendos = [-250, -300, -330]

clientes = [   # nome, setor, vendas 2023–2025
    ("Cliente A", "Instalações elétricas", [2350, 2480, 2620]),
    ("Cliente B", "Construção", [1180, 1210, 1250]),
    ("Cliente C", "Metalomecânica", [960, 1000, 1030]),
    ("Cliente D", "Instalações elétricas", [870, 890, 920]),
    ("Cliente E", "Retalho especializado", [760, 790, 810]),
    ("Cliente F", "Hotelaria", [640, 660, 690]),
    ("Cliente G", "Indústria alimentar", [590, 600, 630]),
    ("Cliente H", "Construção", [520, 540, 550]),
    ("Cliente I", "Climatização (AVAC)", [450, 470, 500]),
    ("Cliente J", "Energia solar", [0, 230, 410]),
]
ajustes = [   # ajustes da gestão ao EBITDA: cada um só num ano
    ("Custos legais de um litígio encerrado", [25, 0, 0]),
    ("Certificação ISO 9001 (projeto único)", [0, 35, 0]),
    ("Migração do sistema informático (ERP)", [0, 0, 40]),
]
covenant = 3.0

# ---------------------------------------------------------------- valores derivados
add = lambda *cols: [sum(v) for v in zip(*cols)]
margem = add(vendas, cmv)
ebitda = add(margem, fse, pess, imp, orend, ogast)
ebit = add(ebitda, da)
rai = add(ebit, juros)
rl = add(rai, irc)

fin_cp = [300 + 60] * 4
fin_lp = [e + l - c for e, l, c in zip(emprestimo, locacoes, fin_cp)]
delta = lambda s: [s[i + 1] - s[i] for i in range(3)]            # variação do balanço 2023–2025
d_cli, d_inv, d_oac = [[-x for x in delta(s)] for s in (cli, inv, oac)]   # ativos: sinal contrário
d_forn, d_opc = delta(forn), delta(opc)
cgo = add(rl, [-x for x in da], [-x for x in juros], [-x for x in irc], [-g for g in ganho_alienacao],
          d_cli, d_inv, d_forn, d_oac, d_opc)
fco = add(cgo, juros, irc)                                         # juros e impostos pagos = gasto do ano
fci = add(capex, alienacao)
pag_fin = [-(x) for x in delta([e + l for e, l in zip(emprestimo, locacoes)])]
fcf_fin = add([-p for p in pag_fin], dividendos)
var_caixa = add(fco, fci, fcf_fin)
caixa = [caixa_2022]
for v in var_caixa:
    caixa.append(caixa[-1] + v)

aft = [aft_2022]
for c, d, a, g in zip(capex, da, alienacao, ganho_alienacao):
    aft.append(aft[-1] - c + d - (a - g))                          # + investimento − depreciação − valor vendido

ativo = add(aft, inv, cli, oac, caixa)
passivo_sem_cp = add(fin_lp, forn, fin_cp, opc)
cp = [ativo[0] - passivo_sem_cp[0]]
for r, dv in zip(rl, dividendos):
    cp.append(cp[-1] + r + dv)
rl_bs = [rl_2022] + rl
reservas = [c - capital_social - r for c, r in zip(cp, rl_bs)]

outros_clientes = [v - sum(c[2][i] for c in clientes) for i, v in enumerate(vendas)]
total_ajustes = [sum(a[1][i] for a in ajustes) for i in range(3)]


def check() -> None:
    """A empresa é 'limpa' por construção: tudo bate entre si."""
    for i in range(4):
        assert ativo[i] == cp[i] + passivo_sem_cp[i], f"balanço não fecha em {BS[i]}"
    for i in range(3):
        assert caixa[i + 1] == caixa[i] + var_caixa[i]
        assert aft[i + 1] > 0 and caixa[i + 1] > 0
    assert reservas[1:] == [reservas[i] + rl_bs[i] + dividendos[i] for i in range(3)]


# ---------------------------------------------------------------- data room
BOLD, TITLE = Font(bold=True), Font(bold=True, size=13)
BLUE = Font(color="1F4E9E")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
NUM = '#,##0;-#,##0;"–"'


def _sheet(wb, title, head, rows, subtitle="€ milhares | Exercícios findos em 31 de dezembro", widths=(52,)):
    """Folha com título, cabeçalho na linha 4 e linhas a partir da 5 (None = linha vazia)."""
    ws = wb.create_sheet(title)
    ws["A1"], ws["A2"] = head[0], subtitle
    ws["A1"].font = TITLE
    for c, h in enumerate(head[1], 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font, cell.fill = BOLD, HEAD_FILL
    for r, row in enumerate(rows, 5):
        if row is None:
            continue
        label, values, *style = row
        ws.cell(row=r, column=1, value=label).font = BOLD if style else Font()
        for c, v in enumerate(values, 2):
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, (int, float)):
                cell.number_format = "0.0%" if isinstance(v, float) and abs(v) < 1.5 and v != int(v) else NUM
                cell.font = BOLD if style else BLUE
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCDEFGHI"[i]].width = w
    for col in "BCDEFGH"[len(widths) - 1:]:
        ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, 12)
    return ws


def write_data_room() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leia-me"
    ws["A1"] = f"{NAME} — data room financeiro"
    ws["A1"].font = TITLE
    ws["A2"] = "EMPRESA FICTÍCIA. Dados criados para formação; qualquer semelhança com empresas reais é coincidência."
    info = [
        ("Contexto", "Um grupo industrial avalia a compra de 100% da Velmora. Este é o data room financeiro "
                     "preliminar que a gestão enviou ao comprador."),
        ("Atividade", "Distribuição B2B de material elétrico (cabo, aparelhagem, iluminação e quadros) a "
                      "instaladores, construtoras e indústria. Sede em Braga."),
        ("Período", "Exercícios de 2023, 2024 e 2025 (ano civil). O balanço inclui 2022 como saldo de abertura."),
        ("Unidades", "Milhares de euros (€ k), salvo indicação em contrário."),
        ("Legenda", "Azul: valores reportados pela empresa. Negrito: totais."),
        ("Auditoria", "Contas de 2023, 2024 e 2025 auditadas, sem reservas."),
    ]
    for r, (k, v) in enumerate(info, 4):
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=v)
    sheets = [("DRE", "Demonstração de resultados 2023–2025"), ("Balanço", "Balanço 2022–2025 (31 de dezembro)"),
              ("DFC", "Demonstração de fluxos de caixa 2023–2025 (método indireto)"),
              ("Vendas por Cliente", "Vendas pelos maiores clientes 2023–2025"),
              ("EBITDA Ajustado", "EBITDA ajustado segundo a gestão (management presentation)"),
              ("Dívida", "Mapa de financiamentos em 31/12/2025"),
              ("Relatório de Gestão", "Excerto do relatório de gestão de 2025")]
    ws["A11"], ws["B11"] = "Folha", "Conteúdo"
    ws["A11"].font = ws["B11"].font = BOLD
    for r, (s, d) in enumerate(sheets, 12):
        ws.cell(row=r, column=1, value=s)
        ws.cell(row=r, column=2, value=d)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 110

    yrs = [str(y) for y in PL]
    _sheet(wb, "DRE", ("Demonstração de resultados", ["Rubrica"] + yrs), [
        ("Vendas e serviços prestados", vendas),
        ("Custo das mercadorias vendidas e das matérias consumidas", cmv),
        ("Margem bruta", margem, "total"),
        ("Fornecimentos e serviços externos", fse),
        ("Gastos com o pessoal", pess),
        ("Imparidade de dívidas a receber (perdas)", imp),
        ("Outros rendimentos", orend),
        ("Outros gastos", ogast),
        ("EBITDA", ebitda, "total"),
        ("Gastos de depreciação e de amortização", da),
        ("Resultado operacional (EBIT)", ebit, "total"),
        ("Juros e gastos similares suportados", juros),
        ("Resultado antes de impostos", rai, "total"),
        ("Imposto sobre o rendimento do período", irc),
        ("Resultado líquido do período", rl, "total"),
        None,
        ("Informação adicional", [], "total"),
        ("Número médio de colaboradores", colaboradores),
    ])
    _sheet(wb, "Balanço", ("Balanço", ["Rubrica"] + [str(y) for y in BS]), [
        ("ATIVO", [], "total"),
        ("Ativos fixos tangíveis", aft),
        ("Total do ativo não corrente", aft, "total"),
        ("Inventários", inv),
        ("Clientes", cli),
        ("Outros ativos correntes", oac),
        ("Caixa e depósitos bancários", caixa),
        ("Total do ativo corrente", add(inv, cli, oac, caixa), "total"),
        ("TOTAL DO ATIVO", ativo, "total"),
        None,
        ("CAPITAL PRÓPRIO E PASSIVO", [], "total"),
        ("Capital social", [capital_social] * 4),
        ("Reservas e resultados transitados", reservas),
        ("Resultado líquido do período", rl_bs),
        ("Total do capital próprio", cp, "total"),
        ("Financiamentos obtidos (não corrente)", fin_lp),
        ("Total do passivo não corrente", fin_lp, "total"),
        ("Fornecedores", forn),
        ("Financiamentos obtidos (corrente)", fin_cp),
        ("Outros passivos correntes", opc),
        ("Total do passivo corrente", add(forn, fin_cp, opc), "total"),
        ("Total do passivo", passivo_sem_cp, "total"),
        ("TOTAL DO CAPITAL PRÓPRIO E PASSIVO", add(cp, passivo_sem_cp), "total"),
    ], subtitle="€ milhares | Em 31 de dezembro")
    neg = lambda s: [-x for x in s]
    _sheet(wb, "DFC", ("Demonstração de fluxos de caixa (método indireto)", ["Rubrica"] + yrs), [
        ("Resultado líquido do período", rl),
        ("Gastos de depreciação e de amortização", neg(da)),
        ("Juros e gastos similares", neg(juros)),
        ("Imposto sobre o rendimento", neg(irc)),
        ("(Ganhos)/perdas na alienação de ativos fixos", neg(ganho_alienacao)),
        ("Variação de clientes", d_cli),
        ("Variação de inventários", d_inv),
        ("Variação de fornecedores", d_forn),
        ("Variação de outros ativos correntes", d_oac),
        ("Variação de outros passivos correntes", d_opc),
        ("Caixa gerada pelas operações", cgo, "total"),
        ("Juros pagos", juros),
        ("Impostos sobre o rendimento pagos", irc),
        ("Fluxo de caixa das atividades operacionais (1)", fco, "total"),
        None,
        ("Pagamentos de ativos fixos tangíveis", capex),
        ("Recebimentos da alienação de ativos fixos", alienacao),
        ("Fluxo de caixa das atividades de investimento (2)", fci, "total"),
        None,
        ("Recebimentos de financiamentos obtidos", [0, 0, 0]),
        ("Pagamentos de financiamentos obtidos", neg(pag_fin)),
        ("Dividendos pagos", dividendos),
        ("Fluxo de caixa das atividades de financiamento (3)", fcf_fin, "total"),
        None,
        ("Variação de caixa (1)+(2)+(3)", var_caixa, "total"),
        ("Caixa no início do período", caixa[:-1]),
        ("Caixa no fim do período", caixa[1:]),
    ])

    ws = _sheet(wb, "Vendas por Cliente", ("Vendas por cliente", ["Cliente", "Setor", "2023", "% 2023", "2024",
                                                                   "% 2024", "2025", "% 2025"]), [],
                subtitle="€ milhares | % do total de vendas do ano", widths=(34, 24))
    rows = [(n, s, v) for n, s, v in clientes] + [("Outros clientes (cerca de 260)", "Vários", outros_clientes)]
    for r, (n, s, v) in enumerate(rows + [("Total de vendas", None, vendas)], 5):
        ws.cell(row=r, column=1, value=n).font = BOLD if s is None else Font()
        ws.cell(row=r, column=2, value=s)
        for i in range(3):
            a = ws.cell(row=r, column=3 + 2 * i, value=v[i])
            b = ws.cell(row=r, column=4 + 2 * i, value=v[i] / vendas[i])
            a.number_format, b.number_format = NUM, "0.0%"
            a.font = BOLD if s is None else BLUE

    _sheet(wb, "EBITDA Ajustado", ("EBITDA ajustado — visão da gestão", ["Rubrica"] + yrs), [
        ("Vendas", vendas),
        ("EBITDA reportado", ebitda),
        ("Ajustamentos (itens não recorrentes, segundo a gestão)", [], "total"),
        *[(n, v) for n, v in ajustes],
        ("Total de ajustamentos", total_ajustes, "total"),
        ("EBITDA ajustado", add(ebitda, total_ajustes), "total"),
        ("Margem EBITDA ajustada", [e / s for e, s in zip(add(ebitda, total_ajustes), vendas)], "total"),
    ], subtitle="€ milhares | Fonte: management presentation, março de 2026")

    ws = wb.create_sheet("Dívida")
    ws["A1"], ws["A2"] = "Mapa de financiamentos em 31/12/2025", "€ milhares"
    ws["A1"].font = TITLE
    head = ["Financiamento", "Tipo de reembolso", "Em dívida 31/12/2024", "Em dívida 31/12/2025",
            "Corrente (≤ 12 meses)", "Não corrente", "Vencimento final", "Taxa de juro", "Garantias e covenants"]
    loans = [
        ["Empréstimo bancário", "Amortizações semestrais", emprestimo[2], emprestimo[3], 300, emprestimo[3] - 300,
         "dez/2029", "Euribor 6M + 1,75%", f"Dívida líquida / EBITDA ≤ {covenant:.1f}x".replace(".", ",")
         + ", testado anualmente sobre as contas anuais"],
        ["Locações financeiras (viaturas e empilhadores)", "Rendas mensais", locacoes[2], locacoes[3], 60,
         locacoes[3] - 60, "2027", "Fixa 4,25%", "—"],
    ]
    total = ["Total", None] + [sum(l[i] for l in loans) for i in (2, 3, 4, 5)]
    for c, h in enumerate(head, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font, cell.fill = BOLD, HEAD_FILL
    for r, row in enumerate(loans + [total], 5):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, int):
                cell.number_format, cell.font = NUM, BOLD if row is total else BLUE
    for col, w in zip("ABCDEFGHI", (44, 30, 14, 14, 14, 14, 14, 20, 60)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Relatório de Gestão")
    ws["A1"], ws["A2"] = "Relatório de gestão 2025 — excerto", "Texto preparado pelo Conselho de Administração"
    ws["A1"].font = TITLE
    gm = [m / v for m, v in zip(margem, vendas)]
    pct = lambda x: f"{x * 100:.1f}".replace(".", ",") + "%"
    mil = lambda x: f"{x / 1000:.1f}".replace(".", ",")
    paragraphs = [
        ("§1 Atividade e vendas",
         f"Em 2025, as vendas atingiram {mil(vendas[2])} milhões de euros, mais {pct(vendas[2] / vendas[1] - 1)} do que "
         "em 2024. O crescimento foi distribuído pela carteira de clientes; o maior cliente representou cerca de 16% "
         "das vendas. O cliente J, do setor da energia solar, começou a comprar em 2024 e reforçou as encomendas em 2025."),
        ("§2 Margens",
         f"A margem bruta desceu ligeiramente, de {pct(gm[1])} para {pct(gm[2])}, porque a subida do preço do cabo de "
         "cobre só foi refletida em parte nos preços de venda."),
        ("§3 Capital de giro",
         "O prazo médio de recebimento manteve-se em cerca de 62 dias e os stocks acompanharam o crescimento da "
         "atividade. A empresa paga aos fornecedores a cerca de 60 dias."),
        ("§4 Investimento e património",
         f"O investimento em equipamentos, viaturas e sistemas atingiu {-capex[2]} mil euros. Foram vendidas duas "
         f"viaturas usadas por {alienacao[2]} mil euros. As instalações de Braga são próprias."),
        ("§5 Financiamento e perspetivas",
         f"A empresa amortizou {pag_fin[2]} mil euros de dívida, conforme o plano contratado, e distribuiu "
         f"{-dividendos[2]} mil euros de dividendos. A dívida líquida / EBITDA ficou em 0,5x, abaixo do limite de "
         "3,0x do contrato de financiamento."),
    ]
    for r, (k, v) in enumerate(paragraphs, 4):
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=v)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 130
    wb.save(OUT_ROOM)


# ---------------------------------------------------------------- template (folha Input)
def input_cells() -> dict[str, dict[str, object]]:
    """Células a escrever no template, com a fonte de cada linha (mesmas linhas do template da Duriva)."""
    cells: dict[str, object] = {}

    def line(row, values, source, cols="CDE"):
        for c, v in zip(cols, values):
            cells[f"{c}{row}"] = v
        cells[f"F{row}"] = source

    for row, values, src in [(6, vendas, "DRE!B5:D5"), (7, cmv, "DRE!B6:D6"), (8, fse, "DRE!B8:D8"),
                             (9, pess, "DRE!B9:D9"), (10, imp, "DRE!B10:D10"), (11, orend, "DRE!B11:D11"),
                             (12, ogast, "DRE!B12:D12"), (13, da, "DRE!B14:D14"), (14, juros, "DRE!B16:D16"),
                             (15, irc, "DRE!B18:D18"), (16, rl, "DRE!B19:D19")]:
        line(row, values, src)
    for row, values, src in [(19, aft, "Balanço!B6:E6"), (20, [0] * 4, "Não existe no data room (0)"),
                             (21, inv, "Balanço!B8:E8"), (22, cli, "Balanço!B9:E9"), (23, oac, "Balanço!B10:E10"),
                             (24, caixa, "Balanço!B11:E11"), (25, cp, "Balanço!B19:E19"),
                             (26, rl_bs, "Balanço!B18:E18"), (27, fin_lp, "Balanço!B20:E20"),
                             (28, [0] * 4, "Não existe no data room (0)"), (29, forn, "Balanço!B22:E22"),
                             (30, fin_cp, "Balanço!B23:E23"), (31, opc, "Balanço!B24:E24")]:
        line(row, values, src, cols="BCDE")
    neg = lambda s: [-x for x in s]
    for row, values, src in [(34, rl, "DFC!B5:D5"), (35, d_cli, "DFC!B10:D10"), (36, d_inv, "DFC!B11:D11"),
                             (37, d_forn, "DFC!B12:D12"), (38, d_oac, "DFC!B13:D13"), (39, d_opc, "DFC!B14:D14"),
                             (40, cgo, "DFC!B15:D15"), (41, fco, "DFC!B18:D18"), (42, capex, "DFC!B20:D20"),
                             (43, alienacao, "DFC!B21:D21"), (44, neg(ganho_alienacao), "DFC!B9:D9"),
                             (45, caixa[:-1], "DFC!B30:D30"), (46, caixa[1:], "DFC!B31:D31")]:
        line(row, values, src)
    for i in range(12):                          # 12 linhas de clientes no template (49–60)
        row = 49 + i
        if i < len(clientes):
            name, _, values = clientes[i]
            cells[f"A{row}"] = name
            line(row, values, f"Vendas por Cliente!C{5 + i}:G{5 + i}")
        else:
            cells.update({f"A{row}": None, f"C{row}": None, f"D{row}": None, f"E{row}": None, f"F{row}": "(vazio)"})
    line(61, outros_clientes, f"Vendas por Cliente!C{5 + len(clientes)}:G{5 + len(clientes)}")
    for i in range(5):                           # 5 linhas de ajustes no template (64–68)
        row = 64 + i
        if i < len(ajustes):
            name, values = ajustes[i]
            cells[f"A{row}"] = name
            line(row, values, f"EBITDA Ajustado!B{8 + i}:D{8 + i}")
        else:
            cells.update({f"A{row}": None, f"C{row}": None, f"D{row}": None, f"E{row}": None, f"F{row}": "(vazio)"})
    cells.update({"E71": emprestimo[3] + locacoes[3], "F71": "Dívida!D7",
                  "E72": fin_cp[3], "F72": "Dívida!E7",
                  "E73": covenant, "F73": "Dívida!I5 (empréstimo bancário, teste anual)"})
    for row in range(76, 86):                    # todos os documentos foram recebidos
        cells[f"B{row}"] = "Sim"
    leia_me = {"A2": "Receitas, margens e geração de caixa, com regras de dívida, capital de giro e consistência. "
                     "Exemplo preenchido: Velmora (empresa fictícia 'limpa', para medir alarmes falsos).",
               "B10": "Empresa desenhada sem problemas plantados, para medir alarmes falsos. Foi criada por quem "
                      "conhece as regras: zero alertas é um teste de sanidade, não uma medida independente."}
    return {"Input": cells, "Leia-me": leia_me}


# O Excel escreve os valores e recalcula as fórmulas (o openpyxl não calcula fórmulas).
PS_FILL = r"""
$ErrorActionPreference = 'Stop'
$cells = Get-Content -Raw -Encoding UTF8 $env:DD_CELLS | ConvertFrom-Json
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false
$xl.DisplayAlerts = $false
try {
    $wb = $xl.Workbooks.Open($env:DD_TEMPLATE)
    foreach ($sheet in $cells.PSObject.Properties) {
        $ws = $wb.Worksheets.Item($sheet.Name)
        foreach ($cell in $sheet.Value.PSObject.Properties) {
            $v = $cell.Value
            if ($null -eq $v) { [void]$ws.Range($cell.Name).ClearContents() }
            elseif ($v -is [string]) { $ws.Range($cell.Name).Value2 = [string]$v }
            else { $ws.Range($cell.Name).Value2 = [double]$v }   # o JSON chega como Decimal; o Excel quer double
        }
    }
    $xl.CalculateFull()
    $wb.Save()
    $wb.Close($false)
} finally {
    $xl.Quit()
    [void][Runtime.InteropServices.Marshal]::ReleaseComObject($xl)
}
"""


def fill_template() -> None:
    shutil.copyfile(TEMPLATE, OUT_TEMPLATE)
    with tempfile.TemporaryDirectory() as tmp:
        cells_file = Path(tmp) / "cells.json"
        cells_file.write_text(json.dumps(input_cells(), ensure_ascii=False), encoding="utf-8")
        env = {**os.environ, "DD_CELLS": str(cells_file), "DD_TEMPLATE": str(OUT_TEMPLATE)}
        subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", PS_FILL],
                       env=env, check=True, timeout=180)


if __name__ == "__main__":
    check()
    write_data_room()
    print(f"Data room: {OUT_ROOM}")
    fill_template()
    print(f"Template:  {OUT_TEMPLATE}")
