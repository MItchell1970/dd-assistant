"""Layer 1b — read the company's own financial statements (no template needed).

Upload the financial data room as it comes from the company and the agent maps each line
to the standard model:
  - Excel: a dictionary of the European statutory formats recognises each line — Portuguese
    SNC, Spanish PGC, French PCG, German HGB, Italian civil code and IFRS in English
    ("Vendas e serviços prestados", "Importe neto de la cifra de negocios", "Chiffre d'affaires",
    "Umsatzerlöse", "Ricavi delle vendite" and "Revenue" are all revenue). The figures are read
    straight from the cells and every line keeps its cell reference as the source. Free, instant
    and fully auditable.
  - PDF: Claude reads the statements (vision) and returns each figure exactly as printed,
    with its page. The code, not the AI, converts units and signs; the consistency rules
    (R12–R17), the reading checks and the review table then check what was read.

Published accounts differ from a model template, so the reader also:
  - works in euros only (the European market): finds the unit (euros, thousands, millions, in any
    of the six languages) and converts the figures to EUR thousands; a file in another currency
    is refused with a clear message rather than converted at a guessed rate;
  - reads European number and date formats (1.234,5 · 31.12.2024 · Exercice 2024);
  - reads income statements by nature (SNC, IFRS) and by function (US GAAP, UK): cost of sales,
    R&D, SG&A, distribution and administrative expenses, with D&A from the cash flow statement;
  - stops at operating profit: lines below it (e.g. 'Other income/(expense), net' in a US P&L)
    are not EBITDA;
  - adds lines the model keeps together (sales + services rendered; capex on tangible +
    intangible assets; inventory write-downs and provisions within EBITDA, as in SNC);
  - reads a label together with its section ('Clientes' under current vs non-current assets);
  - keeps one column per year when quarters and the full year sit side by side;
  - derives the 'other' balance sheet lines from the published totals, so the balance sheet
    check (R12) tests the company's own totals;
  - flags balance sheet lines that may belong in net debt (leases, pensions, tax payable,
    marketable securities…) and lists every line it did not use, for the reviewer.
"""
from __future__ import annotations

import base64
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from pathlib import Path
from typing import BinaryIO, Literal

import anthropic
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from .input_loader import KEY_DOCUMENTS, BalanceItem, CompanyData, Named, load_input


class ExtractionError(Exception):
    """The file could not be read as financial statements."""


def _norm(text) -> str:
    """Lower case, no accents, punctuation as spaces: '(Ganhos)/perdas' -> 'ganhos perdas'."""
    t = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9&%]+", " ", t)).strip()


def _plain(text) -> str:
    """Lower case without accents, keeping symbols (€, $, £): for unit and currency notes."""
    t = unicodedata.normalize("NFKD", str(text)).lower()
    return "".join(ch for ch in t if not unicodedata.combining(ch))


@dataclass(frozen=True)
class Field:
    key: str
    label: str                    # English name in the review table
    statement: str                # pl, bs or cf
    synonyms: tuple[str, ...]     # the main line
    sign: int = 0                 # -1: negative in the model (costs and payments), whatever the file shows
    adds: tuple[str, ...] = ()    # further lines added to it (e.g. services rendered added to sales)
    section: str = ""             # "current" or "noncurrent": rows of the other section are skipped
    optional: bool = False        # not reported as missing (totals, splits, lines often absent)
    zero_if_missing: bool = False
    exclude: str = ""             # regex: labels that start like a synonym but are something else
    patterns: tuple[str, ...] = ()   # regex tried last, for labels too varied for a list
    above_ebit: bool = False      # only lines above operating profit (EBITDA lines)
    adds_when: tuple[str, ...] = ()   # adds only when the main line is one of these (a component, not a total)


def _f(key, label, statement, synonyms, sign=0, adds=(), section="", optional=False, zero=False, exclude="",
       patterns=(), above_ebit=False, adds_when=()) -> Field:
    return Field(key, label, statement, tuple(_norm(s) for s in synonyms), sign, tuple(_norm(s) for s in adds),
                 section, optional or zero, zero, exclude, tuple(patterns), above_ebit,
                 tuple(_norm(s) for s in adds_when))


# Lines by function (cost of sales method — US GAAP, UK, German Umsatzkostenverfahren), in six languages.
_BY_FUNCTION = ["research and development", "research and development expenses", "selling general and administrative",
                "selling general and administrative expenses", "distribution costs", "distribution expenses",
                "selling and distribution expenses", "selling and distribution costs", "selling and marketing",
                "sales and marketing", "marketing and sales", "administrative expenses", "general and administrative",
                "general and administrative expenses", "gastos de distribuicao", "custos de distribuicao",
                "gastos administrativos", "custos administrativos", "gastos de investigacao e desenvolvimento",
                "gastos de distribucion", "gastos comerciales", "gastos de administracion",
                "gastos de investigacion y desarrollo", "frais commerciaux", "couts commerciaux", "frais administratifs",
                "frais generaux et administratifs", "frais de recherche et developpement", "vertriebskosten",
                "allgemeine verwaltungskosten", "verwaltungskosten", "forschungs und entwicklungskosten",
                "costi di distribuzione", "costi commerciali", "spese generali e amministrative", "costi amministrativi",
                "costi di ricerca e sviluppo"]
# Change in manufactured inventories: an income-side line, never flipped with the costs (see _combine).
_PRODUCTION_CHANGE = ["variacao da producao", "variacao nos inventarios da producao", "variacao de inventarios da producao",
                      "change in manufactured inventories", "changes in inventories of finished goods and work in progress",
                      "variacion de existencias de productos terminados y en curso de fabricacion",
                      "variacion de existencias de productos terminados", "production stockee", "bestandsveranderungen",
                      "erhohung oder verminderung des bestands an fertigen und unfertigen erzeugnissen",
                      "variazioni delle rimanenze di prodotti in corso di lavorazione semilavorati e finiti",
                      "variazioni delle rimanenze di prodotti"]
# grand totals ('Summe Aktiva', 'TOTAL ACTIF') are never a current / non-current subtotal, whatever section they close
_GRAND_TOTAL = r"aktiva|passiva|bilanzsumme|^total (?:do |del |de l )?(?:ativo|activo|actif|assets)|^totale attivo|^total passif"
_BEFORE_TAX = r"antes de impostos|antes de impuestos|before tax|avant impot|vor steuern|prima delle imposte"

FIELDS = [
    # ---- income statement (PT SNC, ES PGC, FR PCG, DE HGB, IT OIC, IFRS in English)
    _f("vendas", "Revenue", "pl",
       ["total net sales", "total revenues", "total revenue", "vendas e servicos prestados",
        "vendas e prestacoes de servicos", "volume de negocios", "importe neto de la cifra de negocios",
        "cifra de negocios", "chiffre d affaires net", "montant net du chiffre d affaires", "chiffre d affaires",
        "umsatzerlose", "umsatzerloese", "umsatz", "ricavi delle vendite e delle prestazioni", "ricavi delle vendite",
        "totale ricavi", "ricavi", "fatturato", "netto omzet", "omzet", "vendas", "ventas netas", "ventas", "reditos",
        "redito", "revenue", "revenues", "net revenues", "net revenue", "turnover", "net sales", "sales"],
       adds=["prestacoes de servicos", "prestacao de servicos", "servicos prestados", "rendering of services",
             "services rendered", "prestaciones de servicios"],
       adds_when=["vendas", "ventas", "sales"], above_ebit=True),
    _f("cmv", "Cost of goods sold", "pl",
       ["total cost of sales", "total cost of revenue", "total cost of revenues",
        "custo das mercadorias vendidas e das materias consumidas", "custo das mercadorias vendidas",
        "custo das vendas", "cmvmc", "cmv", "aprovisionamientos", "coste de las ventas", "coste de ventas",
        "consumo de mercaderias", "cout des ventes", "achats consommes", "achats de marchandises", "materialaufwand",
        "herstellungskosten der zur erzielung der umsatzerlose erbrachten leistungen", "umsatzkosten",
        "herstellungskosten", "costi per materie prime sussidiarie di consumo e di merci", "costi per materie prime",
        "per materie prime sussidiarie di consumo e di merci",
        "costo del venduto", "costs of goods sold and materials consumed", "costs of goods sold", "cost of goods sold",
        "cost of materials consumed", "cost of sales", "cost of revenue", "cost of revenues", "cogs"], -1,
       adds=_PRODUCTION_CHANGE + ["variation de stocks de marchandises", "variation de stocks",
                                  "achats de matieres premieres et autres approvisionnements",
                                  "variation de stocks de matieres premieres et autres approvisionnements",
                                  "variazioni delle rimanenze di materie prime sussidiarie di consumo e merci"],
       above_ebit=True),
    _f("gp_pub", "Gross profit (as published)", "pl",
       ["margem bruta", "lucro bruto", "gross profit", "gross margin", "margen bruto", "marge brute", "bruttoergebnis",
        "bruttoergebnis vom umsatz", "margine lordo"], optional=True, exclude=r"%", above_ebit=True),
    _f("fse", "External supplies and services", "pl",
       ["fornecimentos e servicos externos", "fse", "otros gastos de explotacion", "servicios exteriores",
        "autres achats et charges externes", "charges externes", "costi per servizi", "per servizi",
        "external supplies and services",
        "external services", "third party supplies and services", "supplies and services"], -1,
       adds=["costi per godimento di beni di terzi", "per godimento di beni di terzi"],
       adds_when=["costi per servizi", "per servizi"], above_ebit=True),
    _f("pess", "Staff costs", "pl",
       ["gastos com o pessoal", "gastos com pessoal", "custos com o pessoal", "custos com pessoal", "gastos de personal",
        "charges de personnel", "salaires et traitements", "personalaufwand", "costi per il personale", "per il personale",
        "staff costs",
        "personnel costs", "employee costs", "employee benefits expense"], -1,
       adds=["charges sociales"], adds_when=["salaires et traitements"], above_ebit=True),
    _f("opex_fn", "Operating expenses by function (incl. D&A)", "pl", _BY_FUNCTION, -1, adds=_BY_FUNCTION,
       optional=True, above_ebit=True),   # R&D, SG&A, distribution, administrative (P&L by function)
    _f("opex_tot", "Total operating expenses (as published)", "pl", ["total operating expenses"], -1,
       optional=True, above_ebit=True),
    _f("imp", "Impairment of receivables", "pl",
       ["imparidade de dividas a receber", "imparidades de dividas a receber", "imparidade de clientes",
        "perdidas deterioro y variacion de provisiones por operaciones comerciales", "svalutazioni dei crediti",
        "impairment of trade receivables", "impairment of receivables", "expected credit losses", "bad debt expense"],
       -1, optional=True, above_ebit=True),   # often presented within a wider impairment line
    _f("orend", "Other income", "pl",
       ["outros rendimentos e ganhos", "outros rendimentos", "otros ingresos de explotacion", "ingresos accesorios",
        "autres produits d exploitation", "autres produits", "sonstige betriebliche ertrage", "altri ricavi e proventi",
        "other operating income", "other income"],
       adds=["subsidios a exploracao", "operating grants", "subvenciones de explotacion incorporadas al resultado",
             "subvenciones de explotacion", "subventions d exploitation"], above_ebit=True),
    _f("own", "Own work capitalised", "pl",
       ["trabalhos para a propria entidade", "trabajos realizados por la empresa para su activo", "production immobilisee",
        "andere aktivierte eigenleistungen", "aktivierte eigenleistungen",
        "incrementi di immobilizzazioni per lavori interni", "own work capitalised", "own work capitalized",
        "capitalised own work"], optional=True, above_ebit=True),   # costs moved to the balance sheet: a QoE area
    _f("ogast", "Other expenses (incl. write-downs and provisions)", "pl",
       ["outros gastos e perdas", "outros gastos", "otros gastos", "otros resultados", "autres charges",
        "sonstige betriebliche aufwendungen", "oneri diversi di gestione", "other operating expenses", "other expenses",
        "other costs and losses", "other costs"], -1,
       adds=["imparidade de inventarios", "imparidades de inventarios", "perdas por imparidade em inventarios",
             "provisoes", "inventory write downs", "provisions", "impairments of assets", "impairment of assets",
             "impairment losses", "impots taxes et versements assimiles", "accantonamenti per rischi",
             "altri accantonamenti"], above_ebit=True),
    _f("ebitda_pub", "EBITDA (as published)", "pl",
       ["ebitda", "ebitda reportado", "operating profit before depreciation",
        "resultado antes de depreciacoes gastos de financiamento e impostos", "resultado antes de depreciacoes",
        "earnings before interest taxes depreciation and amortisation",
        "earnings before interest taxes depreciation and amortization"], optional=True,
       exclude=r"%|ajustad|adjusted|normali|recorrente|recurring|margem|margin|bereinigt|rettificat|retraite",
       above_ebit=True),
    _f("da", "Depreciation and amortisation", "pl",
       ["gastos de depreciacao e de amortizacao", "gastos reversoes de depreciacao e de amortizacao",
        "depreciacoes e amortizacoes", "depreciacao e amortizacao", "amortizacoes e depreciacoes",
        "amortizacion del inmovilizado", "dotations aux amortissements", "abschreibungen", "ammortamenti e svalutazioni",
        "ammortamento delle immobilizzazioni materiali", "depreciation and amortisation", "depreciation and amortization",
        "depreciation"], -1,
       adds=["ammortamento delle immobilizzazioni immateriali"], adds_when=["ammortamento delle immobilizzazioni materiali"],
       above_ebit=True),
    _f("opinc", "Operating profit (EBIT, as published)", "pl",
       ["resultados operacionais", "resultado operacional", "resultado operacional ebit", "resultado de explotacion",
        "resultat d exploitation", "betriebsergebnis", "ergebnis der betrieblichen tatigkeit",
        "differenza tra valore e costi della produzione", "risultato operativo", "bedrijfsresultaat", "operating profit",
        "operating income", "operating result", "operating results", "profit from operations", "income from operations",
        "results from operating activities", "ebit"], optional=True,
       exclude=r"before depreciation|antes de depreciac|antes de amortiz|ebitda|%|margem|margin|other operating income"),
    _f("juros", "Interest expense", "pl",
       ["juros e gastos similares suportados", "juros e gastos similares", "juros suportados",
        "gastos e perdas financeiros", "gastos e perdas de financiamento", "gastos de financiamento",
        "gastos financeiros", "custos financeiros", "encargos financeiros", "gastos financieros",
        "interets et charges assimilees", "charges financieres", "zinsen und ahnliche aufwendungen",
        "interessi e altri oneri finanziari", "interest expense", "finance costs", "financial costs",
        "financial expenses"], -1),
    _f("pbt", "Profit before tax", "pl",
       ["resultado antes de impostos", "resultados antes de impostos", "resultado antes de impuestos",
        "resultat avant impots", "resultat courant avant impots", "ergebnis vor steuern", "risultato prima delle imposte",
        "profit before tax", "profit before taxation", "income before income taxes",
        "income before provision for income taxes", "earnings before taxes"], optional=True),
    _f("irc", "Income tax", "pl",
       ["imposto sobre o rendimento do periodo", "imposto sobre o rendimento", "impostos sobre o rendimento", "irc",
        "impuesto sobre beneficios", "impots sur les benefices", "steuern vom einkommen und vom ertrag",
        "imposte sul reddito dell esercizio", "imposte sul reddito", "provision for income taxes", "income tax expense",
        "income tax", "income taxes", "tax expense", "taxation"], -1),
    _f("rl_dre", "Net income", "pl",
       ["resultado liquido do periodo", "resultado liquido do exercicio", "resultado liquido consolidado",
        "resultado consolidado liquido", "resultado liquido", "resultado del ejercicio", "resultat net",
        "resultat de l exercice", "benefice ou perte", "jahresuberschuss", "jahresergebnis",
        "utile perdita dell esercizio", "utile dell esercizio", "risultato dell esercizio", "profit after tax",
        "net income", "net profit", "profit for the year", "profit for the period"], exclude=_BEFORE_TAX),
    _f("rl_grupo", "Net income attributable to the group", "pl",
       ["resultado liquido consolidado atribuivel ao grupo", "ao grupo", "atribuivel ao grupo",
        "acionistas da empresa mae", "detentores de capital da empresa mae", "atribuido a la sociedad dominante",
        "part du groupe", "anteil der aktionare", "di pertinenza del gruppo", "owners of the parent",
        "attributable to owners of the parent", "equity holders of the parent",
        "net income attributable to the equity holders", "net income attributable to equity holders",
        "profit attributable to owners"], optional=True),
    # ---- balance sheet
    _f("aft", "Property, plant and equipment", "bs",
       ["ativos fixos tangiveis", "activos fixos tangiveis", "inmovilizado material", "immobilisations corporelles",
        "sachanlagen", "immobilizzazioni materiali", "materiele vaste activa", "property plant and equipment",
        "tangible fixed assets", "tangible assets", "fixed assets"], section="noncurrent"),
    _f("oanc", "Other non-current assets", "bs", ["outros ativos nao correntes", "other non current assets"],
       section="noncurrent", zero=True),
    _f("inv", "Inventories", "bs", ["inventarios", "existencias", "stocks et en cours", "vorrate", "rimanenze",
                                    "voorraden", "inventories", "inventory", "stocks"], section="current"),
    _f("cli", "Trade receivables", "bs",
       ["clientes", "contas a receber de clientes", "clientes por ventas y prestaciones de servicios",
        "deudores comerciales y otras cuentas a cobrar", "creances clients et comptes rattaches",
        "clients et comptes rattaches", "creances clients", "forderungen aus lieferungen und leistungen",
        "crediti verso clienti", "handelsdebiteuren", "trade receivables", "accounts receivable", "trade debtors",
        "trade and other receivables"], section="current"),
    _f("oac", "Other current assets", "bs",
       ["outros ativos correntes", "outras contas a receber", "other current assets", "other receivables"],
       section="current"),
    _f("caixa", "Cash and cash equivalents", "bs",
       ["caixa e depositos bancarios", "caixa e equivalentes de caixa", "caixa e equivalentes a caixa",
        "caixa e seus equivalentes", "caixa e equivalentes", "efectivo y otros activos liquidos equivalentes",
        "disponibilites", "tresorerie et equivalents de tresorerie", "kassenbestand", "flussige mittel",
        "zahlungsmittel und zahlungsmitteläquivalente", "zahlungsmittel", "disponibilita liquide", "liquide middelen",
        "cash and cash equivalents", "cash at bank", "cash"], section="current"),
    _f("cp", "Total equity", "bs",
       ["total do capital proprio", "total shareholders equity", "total stockholders equity", "total patrimonio neto",
        "total capitaux propres", "summe eigenkapital", "totale patrimonio netto", "capital proprio", "patrimonio neto",
        "capitaux propres", "eigenkapital", "patrimonio netto", "eigen vermogen", "total equity", "shareholders equity",
        "equity"], exclude=r"pasivo|passif|liabilities|passivo e"),
    _f("rl_bs", "Net income (balance sheet)", "bs",
       ["resultado liquido do periodo", "resultado liquido do exercicio", "resultado consolidado liquido",
        "resultado liquido consolidado", "resultado liquido", "resultado del ejercicio", "resultat de l exercice",
        "jahresuberschuss", "utile perdita dell esercizio", "net income", "profit for the year"],
       optional=True),   # US and UK balance sheets show retained earnings instead
    _f("fin_lp", "Non-current borrowings", "bs",
       ["financiamentos obtidos nao corrente", "financiamentos obtidos nao correntes", "non current borrowings",
        "long term borrowings", "long term debt", "financiamentos obtidos", "emprestimos obtidos",
        "emprestimos bancarios", "deudas con entidades de credito",
        "emprunts et dettes aupres des etablissements de credit", "emprunts et dettes financieres",
        "verbindlichkeiten gegenuber kreditinstituten", "debiti verso banche", "schulden aan kredietinstellingen",
        "interest bearing loans", "loans and borrowings", "term debt", "borrowings", "bank loans"],
       section="noncurrent"),
    _f("opnc", "Other non-current liabilities", "bs", ["outros passivos nao correntes", "other non current liabilities"],
       section="noncurrent", zero=True),
    _f("forn", "Trade payables", "bs",
       ["fornecedores", "proveedores", "acreedores comerciales y otras cuentas a pagar",
        "dettes fournisseurs et comptes rattaches", "dettes fournisseurs", "fournisseurs",
        "verbindlichkeiten aus lieferungen und leistungen", "debiti verso fornitori", "handelscrediteuren",
        "trade payables", "accounts payable", "trade creditors", "trade and other payables"], section="current"),
    _f("fin_cp", "Current borrowings", "bs",
       ["financiamentos obtidos corrente", "financiamentos obtidos correntes", "current borrowings",
        "short term borrowings", "short term debt", "financiamentos obtidos", "emprestimos obtidos",
        "emprestimos bancarios", "deudas con entidades de credito", "debiti verso banche", "interest bearing loans",
        "loans and borrowings", "term debt", "borrowings", "bank loans", "commercial paper", "bank overdrafts"],
       adds=["commercial paper", "bank overdrafts", "descobertos bancarios", "current portion of long term debt"],
       section="current"),
    _f("opc", "Other current liabilities", "bs",
       ["outros passivos correntes", "outras contas a pagar", "other current liabilities", "other payables"],
       section="current"),
    _f("tot_anc", "Total non-current assets", "bs",
       ["total do ativo nao corrente", "total de ativos nao correntes", "total dos ativos nao correntes",
        "total activo no corriente", "activo no corriente", "total actif non courant", "actif immobilise",
        "total actif immobilise", "anlagevermogen", "totale immobilizzazioni", "totale attivita non correnti",
        "total non current assets", "non current assets"], optional=True, exclude=_GRAND_TOTAL),
    _f("tot_ac", "Total current assets", "bs",
       ["total do ativo corrente", "total de ativos correntes", "total dos ativos correntes", "total activo corriente",
        "activo corriente", "total actif courant", "actif circulant", "total actif circulant", "umlaufvermogen",
        "totale attivo circolante", "attivo circolante", "totale attivita correnti", "total current assets",
        "current assets"], optional=True, exclude=_GRAND_TOTAL),
    _f("tot_pnc", "Total non-current liabilities", "bs",
       ["total do passivo nao corrente", "total de passivos nao correntes", "total dos passivos nao correntes",
        "total pasivo no corriente", "pasivo no corriente", "total passif non courant", "totale passivita non correnti",
        "total non current liabilities", "non current liabilities"], optional=True, exclude=_GRAND_TOTAL),
    _f("tot_pc", "Total current liabilities", "bs",
       ["total do passivo corrente", "total de passivos correntes", "total dos passivos correntes",
        "total pasivo corriente", "pasivo corriente", "total passif courant", "totale passivita correnti",
        "total current liabilities", "current liabilities"], optional=True, exclude=_GRAND_TOTAL),
    _f("tot_at", "Total assets", "bs",
       ["total do ativo", "total ativo", "total del activo", "total activo", "total de l actif", "total actif",
        "summe aktiva", "bilanzsumme", "totale attivo", "totale attivita", "totaal activa", "total assets"],
       optional=True, exclude=r"corrent|courant|circulant|correnti|non"),
    _f("tot_pl", "Total equity and liabilities", "bs",
       ["total do capital proprio e passivo", "total do passivo e capital proprio", "total patrimonio neto y pasivo",
        "total del patrimonio neto y pasivo", "total du passif", "total passif", "summe passiva", "totale passivo",
        "totale passivita e patrimonio netto", "totaal passiva", "total equity and liabilities",
        "total liabilities and equity", "total liabilities and shareholders equity"],
       optional=True, exclude=r"corrent|courant|circulant|correnti|exigible|kurzfristig|langfristig"),
    # ---- cash flow statement
    _f("rl_dfc", "Net income (cash flow)", "cf",
       ["resultado liquido do periodo", "resultado liquido", "resultado del ejercicio", "resultat net",
        "jahresuberschuss", "utile perdita dell esercizio", "net income", "profit for the year"], optional=True,
       exclude=_BEFORE_TAX),
    _f("ganho_alien", "(Gains)/losses on disposals", "cf",
       ["ganhos perdas na alienacao", "perdas ganhos na alienacao", "ganhos na alienacao", "gain loss on disposal",
        "loss gain on disposal", "gain on disposal", "profit on disposal", "resultados por bajas y enajenaciones",
        "plus ou moins values de cession", "plusvalenze minusvalenze derivanti dalla cessione"], optional=True),
    _f("da_cf", "D&A (cash flow statement)", "cf",
       ["depreciation and amortization", "depreciation and amortisation", "depreciation amortization and impairment",
        "depreciation amortisation and impairment", "gastos de depreciacao e de amortizacao",
        "depreciacoes e amortizacoes", "amortizacoes e depreciacoes", "amortizacion del inmovilizado",
        "dotations aux amortissements", "abschreibungen", "ammortamenti delle immobilizzazioni",
        "depreciation of property plant and equipment"],
       adds=["amortisation of intangible assets", "amortization of intangible assets",
             "depreciation of right of use assets"], adds_when=["depreciation of property plant and equipment"],
       optional=True),
    _f("sbc", "Share-based compensation (non-cash)", "cf",
       ["share based compensation expense", "share based compensation", "stock based compensation",
        "share based payments", "pagamentos com base em acoes"], optional=True),
    _f("d_cli", "Change in trade receivables", "cf",
       ["variacao de clientes", "variacao das contas a receber", "change in trade receivables", "change in receivables",
        "increase decrease in trade receivables", "decrease increase in trade receivables"], optional=True),
    _f("d_inv", "Change in inventories", "cf",
       ["variacao de inventarios", "variacao de existencias", "change in inventories",
        "increase decrease in inventories", "decrease increase in inventories"], optional=True),
    _f("d_forn", "Change in trade payables", "cf",
       ["variacao de fornecedores", "change in trade payables", "change in payables",
        "increase decrease in trade payables"], optional=True),
    _f("d_oac", "Change in other current assets", "cf",
       ["variacao de outros ativos correntes", "change in other current assets"], optional=True),
    _f("d_opc", "Change in other current liabilities", "cf",
       ["variacao de outros passivos correntes", "change in other current liabilities"], optional=True),
    # Cash generated from operations, as defined in the indirect method. Direct-method subtotals such as
    # 'Fluxo gerado pelas operações' come before other operating payments and are deliberately not mapped:
    # cash conversion then falls back to the operating cash flow.
    _f("cgo", "Cash generated from operations", "cf",
       ["caixa gerada pelas operacoes", "cash generated from operations"], optional=True),
    _f("fco", "Net cash from operating activities", "cf",
       ["fluxo de caixa das atividades operacionais", "fluxos de caixa das atividades operacionais",
        "fluxo das atividades operacionais", "fluxos das atividades operacionais",
        "flujos de efectivo de las actividades de explotacion", "flux net de tresorerie genere par l activite",
        "flux de tresorerie lies a l activite", "cashflow aus laufender geschaftstatigkeit",
        "cashflow aus der laufenden geschaftstatigkeit", "flusso finanziario dell attivita operativa",
        "flusso finanziario derivante dall attivita operativa", "net cash from operating activities",
        "net cash generated from operating activities", "net cash provided by operating activities",
        "net cash provided by used in operating activities", "net cash from used in operating activities",
        "cash generated by operating activities", "cash flow from operating activities", "operating cash flow"]),
    _f("capex", "Capital expenditure", "cf",
       ["pagamentos de ativos fixos tangiveis", "pagamentos respeitantes a ativos fixos tangiveis",
        "pagamentos de ativos fixos", "pagos por inversiones inmovilizado material",
        "acquisitions d immobilisations corporelles et incorporelles", "acquisitions d immobilisations corporelles",
        "acquisitions d immobilisations", "auszahlungen fur investitionen in das sachanlagevermogen",
        "immobilizzazioni materiali investimenti", "capital expenditure", "capital expenditures",
        "purchase of property plant and equipment", "purchases of property plant and equipment",
        "purchases of property and equipment", "payments for acquisition of property plant and equipment",
        "acquisition of property plant and equipment", "additions to property plant and equipment",
        "payments due to tangible assets", "payments relating to tangible assets", "capex"], -1,
       adds=["pagamentos respeitantes a ativos intangiveis", "pagamentos de ativos intangiveis",
             "pagos por inversiones inmovilizado intangible",
             "auszahlungen fur investitionen in das immaterielle anlagevermogen",
             "immobilizzazioni immateriali investimenti", "acquisitions d immobilisations incorporelles",
             "purchase of intangible assets", "purchases of intangible assets", "payments due to intangible assets",
             "payments relating to intangible assets", "payments for acquisition of intangible assets"]),
    _f("alien", "Proceeds from asset disposals", "cf",
       ["recebimentos da alienacao de ativos fixos", "recebimentos de ativos fixos",
        "recebimentos provenientes de ativos fixos tangiveis", "collections due to tangible assets",
        "cobros por desinversiones inmovilizado material", "cessions d immobilisations",
        "einzahlungen aus abgangen von gegenstanden des sachanlagevermogens", "immobilizzazioni materiali disinvestimenti",
        "proceeds from disposal", "proceeds from sale of property plant and equipment"], optional=True),
    _f("cx_ini", "Cash at beginning of period", "cf",
       ["caixa no inicio do periodo", "caixa e seus equivalentes no inicio do periodo",
        "caixa e equivalentes no inicio do periodo", "cash and cash equivalents at beginning of",
        "cash at beginning of", "opening cash"],
       patterns=(r"^cash\b.*\b(?:beginning|opening)\b", r"^caixa\b.*\binicio\b",
                 r"^efectivo\b.*\b(?:comienzo|inicio)\b", r"\btresorerie\b.*\bouverture\b",
                 r"^finanzmittel\w*\b.*\banfang\b", r"^disponibilita liquide\b.*\binizio\b")),
    _f("cx_fim", "Cash at end of period", "cf",
       ["caixa no fim do periodo", "caixa e equivalentes no fim do periodo", "caixa e seus equivalentes no fim do periodo",
        "cash and cash equivalents at end of", "cash at end of", "closing cash"],
       patterns=(r"^cash\b(?!.*\b(?:beginning|opening)\b).*\b(?:end|ending|closing)\b", r"^caixa\b.*\bfim\b",
                 r"^efectivo\b.*\bfinal\b", r"\btresorerie\b.*\bcloture\b", r"^finanzmittel\w*\b.*\bende\b",
                 r"^disponibilita liquide\b.*\bfine\b")),
]
FIELD_BY_KEY = {f.key: f for f in FIELDS}
_INCOME_LIKE = tuple(_norm(x) for x in _PRODUCTION_CHANGE)
STATEMENTS = {"pl": "Income statement", "bs": "Balance sheet", "cf": "Cash flow statement"}

# 'Other' balance sheet lines derived from the published totals: total less the lines identified.
RESIDUALS = {"oanc": ("tot_anc", ["aft"]), "oac": ("tot_ac", ["inv", "cli", "caixa"]),
             "opnc": ("tot_pnc", ["fin_lp"]), "opc": ("tot_pc", ["forn", "fin_cp"])}

# Balance sheet lines that may belong in net debt. Included by default: a prudent buy-side view (debt-like
# items clearly owed; cash-like items only when current); 'to assess' items are shown but not included.
DEBT_LIKE = [
    ("Lease liabilities", ["passivos de locacao", "passivo de locacao", "responsabilidades por locacoes",
                           "locacoes financeiras", "locacao financeira", "lease liabilities", "lease liability",
                           "finance lease liabilities", "leases", "leasing", "acreedores por arrendamiento financiero",
                           "pasivos por arrendamiento", "dettes de location", "obligations locatives",
                           "leasingverbindlichkeiten", "debiti per leasing", "passivita per leasing"], True),
    ("Post-employment benefit obligations",
     ["responsabilidades por planos de beneficios definidos", "responsabilidades por beneficios pos emprego",
      "beneficios pos emprego", "post employment benefits", "post employment benefit obligations",
      "retirement benefit obligations", "pension obligations", "pension liabilities", "defined benefit obligations",
      "employee benefit obligations", "provisiones por retribuciones a largo plazo al personal",
      "provisions pour retraites", "engagements de retraite", "ruckstellungen fur pensionen", "pensionsruckstellungen",
      "trattamento di fine rapporto di lavoro subordinato", "trattamento di fine rapporto", "fondo tfr"], True),
    ("Income tax payable", ["imposto sobre o rendimento a pagar", "income tax liabilities", "income tax payable",
                            "income taxes payable", "current tax liabilities", "current income tax liabilities",
                            "corporation tax", "pasivos por impuesto corriente",
                            "steuerruckstellungen", "debiti tributari"], True),
    ("Dividends payable", ["dividendos a pagar", "dividends payable"], True),
    ("Shareholder loans", ["suprimentos", "financiamentos de acionistas", "emprestimos de acionistas",
                           "shareholder loans", "loans from shareholders", "deudas con empresas del grupo",
                           "deudas con socios", "comptes courants d associes", "verbindlichkeiten gegenuber gesellschaftern",
                           "debiti verso soci per finanziamenti", "debiti verso controllanti"], True),
    ("Deferred or contingent consideration", ["deferred consideration", "contingent consideration", "earn out"], True),
    ("Factoring and confirming", ["factoring", "confirming", "reverse factoring"], True),
    ("Provisions", ["provisoes", "provisions", "provisoes para riscos e encargos", "provisiones a largo plazo",
                    "provisiones a corto plazo", "provisions pour risques et charges", "sonstige ruckstellungen",
                    "ruckstellungen", "fondi per rischi e oneri"], False),
    ("Other financial liabilities", ["outros passivos financeiros", "other financial liabilities",
                                     "derivative financial instruments", "instrumentos financeiros derivados"], False),
    ("Taxes and social security payable", ["estado e outros entes publicos", "dettes fiscales et sociales",
                                           "otras deudas con las administraciones publicas"], False),
    ("Deferred revenue and customer advances", ["deferred revenue", "contract liabilities", "adiantamentos de clientes",
                                                "customer advances", "advances from customers", "deferred income",
                                                "rendimentos diferidos", "anticipos de clientes",
                                                "avances et acomptes recus", "erhaltene anzahlungen", "acconti"],
     False),
]
CASH_LIKE = [
    ("Marketable securities and short-term investments",
     ["marketable securities", "short term investments", "aplicacoes financeiras", "term deposits",
      "depositos a prazo", "financial assets at fair value through profit or loss",
      "ativos financeiros detidos para negociacao", "inversiones financieras a corto plazo",
      "valeurs mobilieres de placement", "wertpapiere des umlaufvermogens", "sonstige wertpapiere",
      "attivita finanziarie che non costituiscono immobilizzazioni"], "current"),
    ("Other financial assets", ["outros ativos financeiros", "other financial assets"], False),
    ("Restricted cash", ["restricted cash", "caixa restrita", "depositos cativos"], False),
    ("Assets held for sale", ["ativos nao correntes detidos para venda", "non current assets held for sale",
                              "assets held for sale"], False),
]
_DEBT_LIKE = [(c, tuple(_norm(s) for s in syn), inc) for c, syn, inc in DEBT_LIKE]
_CASH_LIKE = [(c, tuple(_norm(s) for s in syn), inc) for c, syn, inc in CASH_LIKE]

# Sheet type from its name (or title), checked in this order.
SHEET_TYPES = [
    ("customers", ["vendas por cliente", "sales by customer", "revenue by customer", "customers", "clientes",
                   "ventas por cliente", "ventes par client", "umsatz nach kunden", "vendite per cliente"]),
    ("adjustments", ["ebitda ajustado", "adjusted ebitda", "normalised ebitda", "normalized ebitda", "ajustamentos",
                     "ebitda retraite", "bereinigtes ebitda", "ebitda rettificato"]),
    ("debt", ["divida", "debt", "financiamentos", "borrowings", "facilities", "loans", "deuda financiera",
              "endettement", "finanzverbindlichkeiten", "debiti finanziari"]),
    ("cf", ["dfc", "fluxos de caixa", "fluxo de caixa", "cash flow", "cash flows", "estado de flujos de efectivo",
            "flujos de efectivo", "tableau des flux de tresorerie", "flux de tresorerie", "kapitalflussrechnung",
            "rendiconto finanziario"]),
    ("bs", ["balanco", "balance sheet", "balance sheets", "posicao financeira", "financial position",
            "balance de situacion", "balance", "bilan", "bilanz", "stato patrimoniale"]),
    ("pl", ["dre", "demonstracao de resultados", "demonstracao dos resultados", "demonstracao consolidada dos resultados",
            "income statement", "profit and loss", "p&l", "p l", "statement of operations", "statements of operations",
            "cuenta de perdidas y ganancias", "perdidas y ganancias", "compte de resultat", "guv",
            "gewinn und verlustrechnung", "conto economico"]),
]
NOT_ADJUSTMENTS = ("vendas", "revenue", "sales", "volume de negocios", "ebitda", "total", "margem", "margin")
OTHER_CUSTOMERS = ("outros", "other", "others", "remaining", "restantes", "demais")


@dataclass
class Extraction:
    """What the agent read: the standard model plus everything needed to review it."""
    data: CompanyData
    company: str
    method: str                                  # how the figures were read
    items: list[dict] = field(default_factory=list)   # review table: one row per standard line
    missing: list[str] = field(default_factory=list)  # standard lines not found
    source_text: str = ""                        # full text of the data room, for the AI memo
    unmapped: list[dict] = field(default_factory=list)   # lines of the statements the analysis does not use
    files: list[dict] = field(default_factory=list)      # the data room's files and how each was read


# ---------------------------------------------------------------- Excel helpers
def _as_year(value) -> int | None:
    """Year in a column header: 2023, '2023', 'FY2023', 'FY23', '2023A', '31/12/2023', 'September 28, 2024'
    or a date."""
    if isinstance(value, (datetime, date)):
        return value.year
    if isinstance(value, (int, float)) and float(value).is_integer() and 1990 <= value <= 2100:
        return int(value)
    if not isinstance(value, str) or "%" in value:
        return None
    t = re.sub(r"\s+", " ", _plain(value).strip())
    m = (re.fullmatch(r"(?:fy\s*)?(\d{4})\s*[aeb]?", t) or re.fullmatch(r"\d{1,2}[/.-]\d{1,2}[/.-](\d{4})", t)
         or re.fullmatch(r"(?:exercicio|ejercicio|exercice|esercizio|geschaftsjahr|gj|ano|year|anno|jahr|annee)"
                         r"\s*(\d{4})", t))
    if m:
        y = int(m.group(1))
        return y if 1990 <= y <= 2100 else None
    m = (re.fullmatch(r"(?:(?:at|a|au|al|am|per|em|en|in|zum|on)\s+)?(?:\d{1,2}\.?\s+)?(?:de\s+)?[a-z]+\.?\s+"
                      r"(?:de\s+)?((?:19|20)\d{2})", t)
         or re.fullmatch(r"[a-z]+\.?\s+\d{1,2},?\s+((?:19|20)\d{2})", t))           # 'Sep. 28, 2024'
    if m:
        return int(m.group(1))
    m = re.fullmatch(r"fy\s*'?(\d{2})", t)
    return 2000 + int(m.group(1)) if m else None


_ANNUAL = re.compile(r"twelve|12 months|year|annual|fiscal|exercicio|anual|\bano\b")
_INTERIM = re.compile(r"three|six|nine|quarter|3 months|6 months|9 months|trimestre|semestre|half")


def _one_column_per_year(ws, header: int, cols: dict[int, int]) -> dict[int, int]:
    """Quarter and full year side by side ('Three Months Ended' / 'Twelve Months Ended'): keep the full-year
    column of each year, from the period header above it; otherwise the rightmost."""
    if len(set(cols.values())) == len(cols):
        return cols

    def period(col: int) -> str:
        for r in range(header - 1, max(header - 4, 0), -1):
            for c in range(col, 0, -1):   # a merged header sits in its first column
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.strip():
                    return _norm(v)
        return ""

    score = lambda col: 2 if _ANNUAL.search(period(col)) else 0 if _INTERIM.search(period(col)) else 1
    best: dict[int, int] = {}
    for col, year in cols.items():
        if year not in best or (score(col), col) >= (score(best[year]), best[year]):
            best[year] = col
    return {col: year for year, col in best.items()}


def _year_header(ws) -> tuple[int | None, dict[int, int]]:
    """Header row and {column number: year}: the first row with at least two year headers."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25)):
        cols = {c.column: y for c in row if (y := _as_year(c.value))}
        if len(cols) >= 2:
            return row[0].row, _one_column_per_year(ws, row[0].row, cols)
    return None, {}


def _num(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):   # "(1,234)", "1 234" or "$ 69,958" written as text; "—" is nil
        t = value.strip()
        for ch in (" ", " ", " ", "R$", "US$", "$", "€", "£"):
            t = t.replace(ch, "")
        if t in ("—", "–", "-"):
            return 0.0
        neg = t.startswith("(") and t.endswith(")")
        t = t.strip("()")
        if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", t) or re.fullmatch(r"-?\d+,\d+", t) and \
                not re.fullmatch(r"-?\d{1,3}(,\d{3})+", t):
            t = t.replace(".", "").replace(",", ".")      # European: 1.234.567,8 · 1234,5
        else:
            t = t.replace(",", "")                        # 1,234,567.8
        if re.fullmatch(r"-?\d+(\.\d+)?", t):
            return -float(t) if neg else float(t)
    return None


def _subtotal_kind(norm: str) -> str:
    """'noncurrent' / 'current' for a subtotal row such as 'Non-current assets' or 'Total do passivo corrente'."""
    t = re.sub(r"^(?:total|totale|summe) (?:do |dos |de |da |das |del |de la |de l |du |des |dell |della )?", "", norm)
    if re.fullmatch(r"(?:non current|ativos? nao correntes?|passivos? nao correntes?|activos? no corrientes?|"
                    r"pasivos? no corrientes?|actifs? non courants?|passifs? non courants?|attivita non correnti|"
                    r"passivita non correnti|anlagevermogen|actif immobilise|immobilizzazioni) ?(?:assets|liabilities)?"
                    r"(?: [a-e])?(?: [a-e])?", t):
        return "noncurrent"
    if re.fullmatch(r"(?:current|ativos? correntes?|passivos? correntes?|activos? corrientes?|pasivos? corrientes?|"
                    r"actifs? courants?|passifs? courants?|attivita correnti|passivita correnti|umlaufvermogen|"
                    r"actif circulant|attivo circolante) ?(?:assets|liabilities)?(?: [a-e])?(?: [a-e])?", t):
        return "current"
    return ""


def _section_kind(section: str) -> str:
    """'noncurrent', 'current' or '' from a section header such as 'Ativos não correntes:'."""
    if re.search(r"nao corrente|non current|no corriente|non courant|non correnti|langfristig|anlagevermogen|"
                 r"actif immobilise|immobilizzazioni", section):
        return "noncurrent"
    return "current" if re.search(r"corrente|current|corriente|courant|correnti|kurzfristig|umlaufvermogen|"
                                  r"actif circulant|attivo circolante", section) else ""


_LIABILITY_HEADER = re.compile(r"capital proprio|passivo|liabilities|equity|pasivo|patrimonio|passif|capitaux propres|"
                               r"passiva|eigenkapital|passivita")
_TOTAL_ASSETS = re.compile(r"(?:total (?:do |dos |del |de l )?(?:ativo|activo|assets|actif)|totale attivo|summe aktiva|"
                           r"bilanzsumme)(?: [a-e])*")   # 'TOTAL ACTIVO (A + B)'
# 'A) ', '1. ', 'IV. ', 'B.6) ' before a label (Spanish PGC, Italian civil code): not part of the name
_ENUMERATION = re.compile(r"^\s*(?:[A-Ha-h]|[IVXivx]{1,4}|\d{1,2})(?:\.\d{1,2})?\s?[).\-–]\s+")


@dataclass
class _Row:
    sheet: str
    row: int
    label: str
    norm: str
    values: dict[int, float | None]
    source: str
    context: str = ""    # section header + label, e.g. 'pagamentos respeitantes a ativos fixos tangiveis'
    kind: str = ""       # current / noncurrent, from the section header
    side: str = "asset"  # asset / liability: balance sheet rows after total assets or an equity header


def _rows(ws, header: int, cols: dict[int, int]) -> list[_Row]:
    """Rows below the header that have a text label and at least one number.

    A label row without numbers ('Passivo corrente', 'Pagamentos respeitantes a:') is a section
    header: the rows under it keep it as context, so a label repeated in two sections can be told apart.
    """
    first, last = min(cols), max(cols)
    out, section, side = [], "", "asset"
    for cells in ws.iter_rows(min_row=header + 1, max_row=ws.max_row):
        label = next((c.value for c in cells if c.column < first and isinstance(c.value, str) and c.value.strip()), None)
        if label is None:
            continue
        values = {cols[c.column]: _num(c.value) for c in cells if c.column in cols}
        if all(v is None for v in values.values()):
            section = _norm(_ENUMERATION.sub("", label))
            side = "liability" if _LIABILITY_HEADER.search(section) else side
            continue
        r = cells[0].row
        source = f"{ws.title}!{get_column_letter(first)}{r}:{get_column_letter(last)}{r}"
        norm = _norm(_ENUMERATION.sub("", label))
        out.append(_Row(ws.title, r, label.strip(), norm, values, source, f"{section} {norm}".strip(),
                        _section_kind(section), side))
        side = "liability" if _TOTAL_ASSETS.fullmatch(norm) else side
    # no section headers? subtotals give their kind. Most formats put the subtotal after its lines
    # ('Current assets 811,818' closes the block); the Spanish PGC puts it first ('A) ACTIVO NO CORRIENTE 5.000'
    # opens it). A balance sheet that starts with a subtotal is read top-down.
    if out and _subtotal_kind(out[0].norm):
        kind = ""
        for row in out:
            kind = _subtotal_kind(row.norm) or ("" if _TOTAL_ASSETS.fullmatch(row.norm) else kind)
            row.kind = row.kind or kind
        return out
    start = 0
    for i, row in enumerate(out):
        if kind := _subtotal_kind(row.norm):
            for prior in out[start:i + 1]:
                prior.kind = prior.kind or kind
            start = i + 1
    return out


def _sheet_type(ws) -> str | None:
    names = [_norm(ws.title), _norm(ws["A1"].value or "")]
    for kind, hints in SHEET_TYPES:
        for name in names:
            if any(re.search(rf"(^| ){re.escape(h)}($| )", name) for h in hints):
                return kind
    return None


def _find(rows: list[_Row], synonyms, used: set, section: str = "", exclude: str = "",
          patterns: tuple[str, ...] = ()) -> list[_Row]:
    """Rows matching the synonyms, best first: exact matches, then labels starting with a synonym;
    each tried on the label alone and on section header + label; then the regex patterns, if any."""
    ok = lambda r: ((r.sheet, r.row) not in used and not (section and r.kind and r.kind != section)
                    and not (exclude and re.search(exclude, r.norm)))
    found = []
    for exact in (True, False):
        for syn in synonyms:
            for r in rows:
                if r in found or not ok(r):
                    continue
                if any(text == syn if exact else text.startswith(syn + " ") for text in (r.norm, r.context)):
                    found.append(r)
    for p in patterns:
        found += [r for r in rows if r not in found and ok(r) and re.search(p, r.norm)]
    return found


def _combine(parts: dict[str, list[tuple[str, dict[int, float | None]]]]) -> dict[str, dict[int, float | None]]:
    """Add up the rows of each line, with costs and payments negative in the model.

    The sign convention is decided for the whole income statement: if its main costs (COGS, external
    services, staff, expenses by function) are shown as positive numbers, every cost row is flipped —
    except a change in manufactured inventories, an income-side row added to cost of sales; otherwise each
    row keeps its published sign, so a net reversal (e.g. of a receivables impairment) stays a gain. A
    cash flow line is flipped only when all its values are positive.
    """
    neg = lambda vals: {y: None if v is None else -v for y, v in vals.items()}
    main_costs = [v for k in ("cmv", "fse", "pess", "opex_fn") for _, vals in parts.get(k, [])[:1]
                  for v in vals.values() if v]
    costs_positive = bool(main_costs) and sum(main_costs) > 0
    out = {}
    for key, rows in parts.items():
        f = FIELD_BY_KEY.get(key)
        total: dict[int, float | None] = {}
        if f and f.statement == "pl":
            for norm, vals in rows:
                flip = f.sign < 0 and costs_positive and not norm.startswith(_INCOME_LIKE)
                total = _add(total, neg(vals) if flip else vals)
        else:
            for _, vals in rows:
                total = _add(total, vals)
            nonzero = [v for v in total.values() if v]
            if f and f.sign < 0 and nonzero and all(v > 0 for v in nonzero):
                total = neg(total)
        out[key] = total
    return out


def _add(a: dict[int, float | None], b: dict[int, float | None]) -> dict[int, float | None]:
    years = set(a) | set(b)
    return {y: None if a.get(y) is None and b.get(y) is None else (a.get(y) or 0.0) + (b.get(y) or 0.0)
            for y in years}


def _derive_residuals(lines: dict, sources: dict, found_as: dict) -> set[str]:
    """Other assets and liabilities = published total less the lines identified (R12 then tests real totals)."""
    derived = set()
    for key, (total, parts) in RESIDUALS.items():
        if total not in lines:
            continue
        lines[key] = {y: None if t is None else t - sum(lines.get(p, {}).get(y) or 0.0 for p in parts)
                      for y, t in lines[total].items()}
        sources[key] = f"{sources[total]} less the lines identified"
        found_as[key] = f"Derived: {FIELD_BY_KEY[total].label.lower()} less identified lines"
        derived.add(key)
    # statutory formats (French, German, Italian) often split neither assets nor liabilities by maturity:
    # what the lines identified do not explain is taken from total assets / total equity and liabilities
    for key, total, parts in (("oac", "tot_at", ["aft", "oanc", "inv", "cli", "caixa"]),
                              ("opc", "tot_pl", ["cp", "fin_lp", "opnc", "forn", "fin_cp"])):
        if key in derived or total not in lines:
            continue
        lines[key] = {y: None if t is None else t - sum(lines.get(p, {}).get(y) or 0.0 for p in parts)
                      for y, t in lines[total].items()}
        sources[key] = f"{sources[total]} less the lines identified"
        found_as[key] = (f"Derived: {FIELD_BY_KEY[total].label.lower()} less identified lines (no current / "
                         f"non-current subtotals: may include non-current items — review before using in NWC)")
        derived.add(key + "_grand")
    return derived


def _standardise(lines: dict, sources: dict, found_as: dict) -> tuple[str, set[str]]:
    """Bring any presentation to the model: P&L by function, residual 'other' lines, zero-filled lines.

    By function (cost of sales, R&D, SG&A…), D&A sits inside the functions, so it is taken from the cash
    flow statement and EBITDA becomes operating profit + D&A, as an analyst would compute it.
    """
    fmt = "nature"
    if not ("fse" in lines or "pess" in lines) and ("opex_fn" in lines or "opex_tot" in lines):
        fmt = "function"
        if "opex_fn" not in lines:
            lines["opex_fn"], sources["opex_fn"] = dict(lines["opex_tot"]), sources["opex_tot"]
            found_as["opex_fn"] = found_as["opex_tot"]
        if "da" not in lines and "da_cf" in lines:
            lines["da"] = {y: None if v is None else -abs(v) for y, v in lines["da_cf"].items()}
            sources["da"], found_as["da"] = sources["da_cf"], f"{found_as['da_cf']} (cash flow statement)"
    derived = _derive_residuals(lines, sources, found_as)
    bs_years = sorted({y for k in ("cli", "cp", "caixa") for y, v in lines.get(k, {}).items() if v is not None})
    for f in FIELDS:
        if f.key not in lines and f.zero_if_missing:
            lines[f.key] = {y: 0.0 for y in bs_years}
            sources[f.key], found_as[f.key] = "Not in data room (0)", "—"
    return fmt, derived


def _classify(norm: str, side: str) -> tuple[str, str, bool | str] | None:
    """(category, 'debt' or 'cash', included by default) for a balance sheet line that may belong in net debt."""
    for nature, table in (("debt", _DEBT_LIKE), ("cash", _CASH_LIKE)):
        if (nature == "debt") != (side == "liability"):
            continue
        for category, synonyms, include in table:
            if any(norm == s or norm.startswith(s + " ") for s in synonyms):
                return category, nature, include
    return None


def _balance_item(category, nature, include, label, kind, values, source, derived) -> BalanceItem:
    other = "opc" if nature == "debt" else "oac"
    # inside the 'other' line when it is derived from a total: its own section, or the grand total (any section)
    in_nwc = (kind == "current" and other in derived) or f"{other}_grand" in derived
    return BalanceItem(category, label, nature, kind, values, source,
                       include if isinstance(include, bool) else kind == include, in_nwc)


# Unit and currency notes, e.g. 'Montantes expressos em Euros', '€ milhares', '(In millions, except …)'.
_UNIT_NOTE = re.compile(r"values|amounts|montantes|valores|expressos|unidade|unit\b|in thousand|in million|"
                        r"em milhares|em milhoes|\bin millions\b|importes|cifras|montants|chiffres|angaben|betrage|"
                        r"importi|valori|bedragen|miles de|milliers|tausend|migliaia")
_UNITS = [("thousands", 1.0, r"milhares|\bmiles\b|milliers|migliaia|duizend|thousands?|tausend|\btsd\b|\bteur\b|"
                             r"t ?€|'000|€ ?k\b|k ?€|\beur ?k\b|\bk ?eur\b|x ?1[ .]?000\b"),
          ("millions", 1000.0, r"milhoes|millones|millions?|milioni|miljoen|\bmio\b|\bmln\b|\bmeur\b|m ?€|€ ?m\b|"
                               r"\beur ?m\b|\bmn\b"),
          ("units", 0.001, r"\beuros?\b")]
# Currencies other than the euro are recognised only to refuse the file ('reais' and 'pounds' are left out on
# purpose: 'empresas reais' means real companies, 'pounds of cork' is a weight).
_CURRENCIES = [("BRL", r"r\$|\bbrl\b"), ("USD", r"us\$|\busd\b|\$|\bdollars?\b|\bdolares\b"),
               ("GBP", r"£|\bgbp\b|sterling"), ("CHF", r"\bchf\b"), ("SEK", r"\bsek\b"), ("NOK", r"\bnok\b"),
               ("DKK", r"\bdkk\b"), ("PLN", r"\bpln\b"), ("EUR", r"€|\beur\b|\beuros?\b|\bteur\b|\bmeur\b")]


def _currency_of(text: str) -> str | None:
    t = _plain(text)
    return next((cur for cur, p in _CURRENCIES if re.search(p, t)), None)


def _unit(wb) -> tuple[float, str, str]:
    """Scale to EUR thousands, how the figures were converted, and the currency (always EUR).

    Unit: the first unit word of a short note in the first rows of a sheet, then of a unit note anywhere
    ('Values in thousand euros.', 'Miles de euros', 'TEUR'); in '(In millions, except shares, in thousands)'
    millions comes first. Currency: from the same notes, else the symbol most used in the cells; EUR if none.
    A file in another currency raises ExtractionError: the analysis is built for euro statements.
    """
    top = [c.value for ws in wb.worksheets for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8))
           for c in row if isinstance(c.value, str)]
    anywhere = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and _UNIT_NOTE.search(_plain(c.value))]
    unit = scale = currency = None
    for text in top + anywhere:
        t = _plain(text)
        if len(text) > 160 or (len(text) > 80 and not _UNIT_NOTE.search(t)):   # a sentence, not a unit note
            continue
        if unit is None:
            hits = [(m.start(), name, s) for name, s, p in _UNITS if (m := re.search(p, t))]
            if hits:
                _, unit, scale = min(hits)
                currency = currency or _currency_of(text)
        if unit and currency:
            break
    if currency is None:
        counts = Counter(cur for ws in wb.worksheets for row in ws.iter_rows() for c in row
                         if isinstance(c.value, str) and len(c.value) <= 12 and (cur := _currency_of(c.value)))
        currency = counts.most_common(1)[0][0] if counts else None
    if currency not in (None, "EUR"):
        raise ExtractionError(f"The figures in this file appear to be in {currency}. DD Assistant analyses statements "
                              "in euros (European market): convert them to euros first, or check the unit note.")
    text = {"thousands": "EUR thousands", "millions": "EUR millions, converted to EUR thousands",
            "units": "euros, converted to EUR thousands", None: "EUR thousands (no unit stated in the file)"}[unit]
    if currency is None:
        text += "; no currency stated, EUR assumed"
    return scale or 1.0, text, "EUR"


def _company(wb, name: str) -> str:
    """Company name from the first sheet's title, or from the file name when the title is a statement."""
    title = wb.worksheets[0]["A1"].value
    if isinstance(title, str) and title.strip() and not _sheet_type(wb.worksheets[0]) and \
            not re.search(r"^(?:consolidated |individual |condensed )?(?:demonstra|balan|statement|balance)",
                          title.lower()):
        return re.split(r"\s+[—–-]\s+|:", title.strip())[0].strip()
    return re.sub(r"[_-]+", " ", Path(name).stem).strip() if name else "Company"


def _covenant(wb, debt_sheet) -> tuple[float | None, str]:
    """Maximum net debt / EBITDA written in the debt schedule (e.g. '≤ 3,0x')."""
    sheets = [debt_sheet] if debt_sheet is not None else []
    for ws in sheets + [s for s in wb.worksheets if s is not debt_sheet]:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "ebitda" in c.value.lower():
                    m = re.search(r"(\d+(?:[.,]\d+)?)\s*x\b", c.value)
                    if m:
                        return float(m.group(1).replace(",", ".")), f"{ws.title}!{c.coordinate}"
    return None, ""


def _debt_schedule(ws) -> tuple[float | None, float | None, str, str]:
    """Total and current portion from the debt schedule's total row, at the latest date."""
    total_col = header = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15)):
        years, current_col = {}, None
        for c in row:
            n = _norm(c.value or "")
            if not n:
                continue
            if re.match(r"(corrente|current|curto prazo|short term|due within)", n):
                current_col = c.column
            elif (m := re.search(r"(20\d{2})", n)):
                years[c.column] = int(m.group(1))
        if current_col and years:
            total_col, header = max(years, key=years.get), row[0].row
            break
    if not header:
        return None, None, "", ""
    for cells in ws.iter_rows(min_row=header + 1, max_row=ws.max_row):
        if _norm(cells[0].value or "").startswith("total"):
            r = cells[0].row
            tot, cur = _num(ws.cell(r, total_col).value), _num(ws.cell(r, current_col).value)
            return (tot, cur, f"{ws.title}!{get_column_letter(total_col)}{r}",
                    f"{ws.title}!{get_column_letter(current_col)}{r}")
    return None, None, "", ""


def data_room_text(path: str | BinaryIO) -> str:
    """Every sheet of an Excel data room as text, with each cell reference (e.g. B5: 18400)."""
    from .report import fmt_value
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out += ["", f"### Sheet '{ws.title}'"]
        for row in ws.iter_rows():
            cells = [f"{c.coordinate}: {fmt_value(c.value)}" for c in row if c.value not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


# ---------------------------------------------------------------- Excel
def extract_excel(file: str | BinaryIO, name: str = "") -> Extraction:
    """Map the company's own Excel financial statements to the standard model."""
    wb = load_workbook(file, data_only=True)
    sheets = {}   # type -> list of (worksheet, header row, year columns)
    for ws in wb.worksheets:
        header, cols = _year_header(ws)
        kind = _sheet_type(ws)
        if kind:
            sheets.setdefault(kind, []).append((ws, header, cols))
        elif cols:
            sheets.setdefault("other", []).append((ws, header, cols))
    scale, unit, currency = _unit(wb)

    parts_by_key, sources, found_as, used, all_rows = {}, {}, {}, set(), {}
    for st in STATEMENTS:
        candidates = [s for s in sheets.get(st, []) if s[2]] or sheets.get("other", [])
        rows = all_rows[st] = [r for ws, header, cols in candidates for r in _rows(ws, header, cols)]
        # lines below operating profit (financial results, 'other income/(expense), net') are not EBITDA
        opinc = FIELD_BY_KEY["opinc"]
        ebit = _find(rows, opinc.synonyms, set(), exclude=opinc.exclude)[:1] if st == "pl" else []
        above = [r for r in rows if not (ebit and r.sheet == ebit[0].sheet and r.row > ebit[0].row)]
        for f in (f for f in FIELDS if f.statement == st):
            pool = above if f.above_ebit else rows
            main = _find(pool, f.synonyms, used, f.section, f.exclude, f.patterns)[:1]
            if not main:
                continue
            parts = main
            if f.adds and (not f.adds_when or main[0].norm in f.adds_when):
                parts = main + _find(pool, f.adds, used | {(main[0].sheet, main[0].row)}, f.section, f.exclude)
            for r in parts:
                used.add((r.sheet, r.row))
            parts_by_key[f.key] = [(r.norm, r.values) for r in parts]
            sources[f.key] = " + ".join(r.source for r in parts)
            found_as[f.key] = " + ".join(r.label for r in parts)
    if "vendas" not in parts_by_key:
        raise ExtractionError("No revenue line found. Is this a financial statements file (P&L, balance sheet, "
                              "cash flow)?")
    lines = _combine(parts_by_key)
    scaled = lambda vals: {y: None if v is None else v * scale for y, v in vals.items()}
    if scale != 1.0:
        lines = {k: scaled(vals) for k, vals in lines.items()}
    pl_format, derived = _standardise(lines, sources, found_as)
    pl_years = sorted(y for y, v in lines["vendas"].items() if v is not None)
    bs_years = sorted({y for k in ("cli", "cp", "caixa") for y, v in lines.get(k, {}).items() if v is not None})
    years = sorted(set(pl_years) | set(bs_years))
    latest = pl_years[-1]

    data = CompanyData(years=years, lines=lines, sources=sources, currency=currency, pl_format=pl_format)
    unmapped = []
    for st, rows in all_rows.items():
        for r in rows:
            if (r.sheet, r.row) in used:
                continue
            hit = _classify(r.norm, r.side) if st == "bs" else None
            if hit:
                data.balance_items.append(_balance_item(*hit, r.label, r.kind, scaled(r.values), r.source, derived))
            else:
                unmapped.append({"Statement": STATEMENTS[st], "Line in the file": r.label, "Source": r.source,
                                 **{f"FY{str(y)[2:]}": v for y, v in sorted(scaled(r.values).items())}})

    for ws, header, cols in sheets.get("customers", []):
        if not cols:
            continue
        for r in _rows(ws, header, cols):
            if r.norm.startswith("total"):
                continue
            if r.norm.startswith(OTHER_CUSTOMERS):
                data.other_clients = scaled(r.values)
            else:
                data.clients.append(Named(r.label, scaled(r.values)))
                data.clients_source = data.clients_source or r.source
        break
    for ws, header, cols in sheets.get("adjustments", []):
        if not cols:
            continue
        for r in _rows(ws, header, cols):
            if not r.norm.startswith(NOT_ADJUSTMENTS):
                data.adjustments.append(Named(r.label, scaled(r.values)))
                data.adjustments_source = data.adjustments_source or r.source
        break

    debt_ws = sheets["debt"][0][0] if sheets.get("debt") else None
    if debt_ws is not None:
        tot, cur, s_tot, s_cur = _debt_schedule(debt_ws)
        if tot is not None:
            lines["map_tot"], sources["map_tot"], found_as["map_tot"] = {latest: tot * scale}, s_tot, "Total"
            lines["map_cur"], sources["map_cur"], found_as["map_cur"] = \
                {latest: None if cur is None else cur * scale}, s_cur, "Total, current"
    cov, s_cov = _covenant(wb, debt_ws)
    if cov is not None:
        lines["cov_lim"], sources["cov_lim"], found_as["cov_lim"] = {latest: cov}, s_cov, f"{cov:g}x"

    sheet_names = " | ".join(_norm(ws.title) for ws in wb.worksheets)
    ageing = lambda *words: bool(re.search(r"aging|ageing|antiguidade", sheet_names)) and any(w in sheet_names for w in words)
    received = {
        "Income statement": True, "Balance sheet": "cli" in lines, "Cash flow statement": "fco" in lines,
        "Sales by customer": bool(data.clients), "Debt schedule": "map_tot" in lines,
        "Support for EBITDA adjustments (invoices, contracts)": False,
        "Breakdown of other income and expenses": any(w in sheet_names for w in ("outros rendimentos", "other income")),
        "AR ageing": ageing("client", "receiv", "devedor"), "Inventory ageing": ageing("invent", "stock", "existenc"),
        "Facility agreements": False,
    }
    data.documents = [(doc, received[doc]) for doc in KEY_DOCUMENTS]

    if hasattr(file, "seek"):
        file.seek(0)
    how = "P&L by function, EBITDA = operating profit + D&A" if pl_format == "function" else "P&L by nature"
    ext = Extraction(data, _company(wb, name), f"Excel — dictionary mapping, figures read from the cells ({unit}; "
                                               f"{how})", source_text=data_room_text(file), unmapped=unmapped)
    _fill_review(ext, found_as)
    return ext


def _fill_review(ext: Extraction, found_as: dict[str, str]) -> None:
    """Review table: every standard line, what it was matched to, its source and its values."""
    d = ext.data
    extra = [("map_tot", "Debt schedule — total outstanding", "Debt"),
             ("map_cur", "Debt schedule — due within 12 months", "Debt"),
             ("cov_lim", "Covenant — maximum net debt / EBITDA", "Debt")]
    # a P&L by function has no lines by nature (external services, staff, other income/expenses)
    by_nature = {"fse", "pess", "orend", "ogast"} if d.pl_format == "function" else set()
    rows = [(f.key, f.label, STATEMENTS[f.statement], f.optional or f.key in by_nature) for f in FIELDS]
    rows += [r + (False,) for r in extra]
    for key, label, statement, optional in rows:
        values = d.lines.get(key)
        item = {"key": key, "Statement": statement, "Line item": label, "Found as": found_as.get(key, "—"),
                "Source": d.sources.get(key, "—")}
        item.update({f"FY{str(y)[2:]}": (values or {}).get(y) for y in d.years})
        ext.items.append(item)
        if values is None and not optional:
            ext.missing.append(label)


def apply_edits(data: CompanyData, edits: dict[str, dict[int, float | None]]) -> CompanyData:
    """A copy of the data with the reviewer's corrections to some lines."""
    lines = {k: dict(v) for k, v in data.lines.items()}
    for key, values in edits.items():
        lines.setdefault(key, {}).update(values)
    return replace(data, lines=lines)


# ---------------------------------------------------------------- PDF (Claude reads it)
MODEL = "claude-opus-5-5"
_KEYS = tuple(f.key for f in FIELDS)


class _YearValue(BaseModel):
    year: int
    value: float


class _Line(BaseModel):
    key: Literal[_KEYS]
    label_in_document: str
    page: int
    values: list[_YearValue]


class _NamedLine(BaseModel):
    name: str
    page: int
    values: list[_YearValue]


class _BalanceLine(BaseModel):
    label_in_document: str
    page: int
    side: Literal["asset", "liability"]
    current: bool
    values: list[_YearValue]


class _Excerpt(BaseModel):
    title: str
    page: int
    text: str


class _PdfStatements(BaseModel):
    company_name: str
    currency: Literal["EUR", "other"]
    unit: Literal["units", "thousands", "millions"]
    line_items: list[_Line]
    other_balance_sheet_lines: list[_BalanceLine]
    customers: list[_NamedLine]
    other_customers: list[_YearValue]
    adjustments: list[_NamedLine]
    debt_total_latest_year: float | None
    debt_current_portion_latest_year: float | None
    covenant_max_net_debt_to_ebitda: float | None
    narrative_excerpts: list[_Excerpt]


_PDF_PROMPT = """You are reading a company's financial statements for a buy-side financial due diligence model.

Return, for every line item below that the document presents, the figures EXACTLY as printed: no calculations,
no rounding, no sign changes, no unit conversion. Give the page where you read each item and the label used in
the document. Omit items the document does not present. Use one entry per fiscal year shown (full years only,
not quarters). When the document splits an item across several lines (e.g. sales and services rendered; capex on
tangible and intangible assets; research and development and SG&A in an income statement by function), return
each line separately with the same key: the code adds them up. Use the consolidated statements when both
consolidated and individual accounts are presented. Lines below operating profit are never EBITDA lines.

Line items (key: meaning):
{fields}

List every other balance sheet line not covered by a key (label as printed, asset or liability, current or not):
the code looks for debt-like and cash-like items among them.

Also return, when presented: sales by customer (named customers and the 'other customers' line), management's
EBITDA adjustments (one entry per adjustment, as positive add-backs if presented that way), the debt schedule's
total and current portion for the latest year, and the maximum net debt / EBITDA covenant.

Finally, copy verbatim (do not paraphrase) the narrative passages a buyer would need: management report,
outlook, financing and refinancing, covenants, going concern, auditor's opinion, subsequent events, related
parties. At most 15 excerpts.

State the currency ('EUR' or 'other') and the unit of the figures (units, thousands or millions)."""


def extract_pdf(pdf: bytes, name: str = "", client: anthropic.Anthropic | None = None) -> Extraction:
    """Claude reads the PDF statements; the code converts units and signs and builds the standard model."""
    client = client or anthropic.Anthropic()
    fields = "\n".join(f"- {f.key}: {f.label} ({STATEMENTS[f.statement]})" for f in FIELDS)
    try:
        response = client.beta.messages.parse(
            model=MODEL,
            max_tokens=16000,
            thinking={"type": "adaptive"},
            output_config={"effort": "high"},
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",     # if the model declines, the API retries on the recommended fallback model
            output_format=_PdfStatements,
            messages=[{"role": "user", "content": [
                {"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                "data": base64.standard_b64encode(pdf).decode()}},
                {"type": "text", "text": _PDF_PROMPT.format(fields=fields)},
            ]}],
        )
    except TypeError as e:
        if "authentication method" in str(e):
            raise ExtractionError("Reading PDFs uses the Claude API: add ANTHROPIC_API_KEY to the .env file.") from e
        raise
    if response.stop_reason == "refusal" or response.parsed_output is None:
        raise ExtractionError("Claude could not read this PDF (refused or truncated). Try the Excel version.")
    return _from_pdf(response.parsed_output, name)


def _from_pdf(p: _PdfStatements, name: str) -> Extraction:
    if p.currency != "EUR":
        raise ExtractionError("The statements in this PDF are not in euros. DD Assistant analyses statements in "
                              "euros (European market).")
    scale = {"units": 0.001, "thousands": 1.0, "millions": 1000.0}[p.unit]   # the model works in EUR thousands
    currency = "EUR"
    vals = lambda yv: {v.year: v.value * scale for v in yv}
    parts, sources, found_as = {}, {}, {}
    for line in p.line_items:   # lines with the same key are added up (sales + services, ...)
        parts.setdefault(line.key, []).append((_norm(line.label_in_document), {v.year: v.value for v in line.values}))
        sources[line.key] = " + ".join(filter(None, [sources.get(line.key), f"{name} p.{line.page}"]))
        found_as[line.key] = " + ".join(filter(None, [found_as.get(line.key), line.label_in_document]))
    lines = {k: {y: None if v is None else v * scale for y, v in vals_.items()} for k, vals_ in _combine(parts).items()}
    if "vendas" not in lines:
        raise ExtractionError("No revenue line found in the PDF.")
    pl_format, derived = _standardise(lines, sources, found_as)
    pl_years = sorted(lines["vendas"])
    bs_years = sorted({y for k in ("cli", "cp", "caixa") for y in lines.get(k, {})})
    latest = pl_years[-1]
    for key, value, label in [("map_tot", p.debt_total_latest_year, "Debt schedule total"),
                              ("map_cur", p.debt_current_portion_latest_year, "Debt schedule current"),
                              ("cov_lim", p.covenant_max_net_debt_to_ebitda, "Covenant")]:
        if value is not None:
            lines[key] = {latest: value if key == "cov_lim" else value * scale}
            sources[key], found_as[key] = f"{name}", label

    data = CompanyData(years=sorted(set(pl_years) | set(bs_years)), lines=lines, sources=sources, currency=currency,
                       pl_format=pl_format)
    unmapped = []
    for b in p.other_balance_sheet_lines:
        side, kind = b.side, "current" if b.current else "noncurrent"
        hit = _classify(_norm(b.label_in_document), side)
        if hit:
            data.balance_items.append(_balance_item(*hit, b.label_in_document, kind, vals(b.values),
                                                    f"{name} p.{b.page}", derived))
        else:
            unmapped.append({"Statement": "Balance sheet", "Line in the file": b.label_in_document,
                             "Source": f"{name} p.{b.page}", **{f"FY{str(y)[2:]}": v for y, v in vals(b.values).items()}})
    data.clients = [Named(c.name, vals(c.values)) for c in p.customers]
    data.clients_source = f"{name} p.{p.customers[0].page}" if p.customers else ""
    data.other_clients = vals(p.other_customers)
    data.adjustments = [Named(a.name, vals(a.values)) for a in p.adjustments]
    data.adjustments_source = f"{name} p.{p.adjustments[0].page}" if p.adjustments else ""
    received = {"Income statement": True, "Balance sheet": "cli" in lines, "Cash flow statement": "fco" in lines,
                "Sales by customer": bool(p.customers), "Debt schedule": "map_tot" in lines}
    data.documents = [(doc, received.get(doc, False)) for doc in KEY_DOCUMENTS]

    text = "\n".join(f"[{name} p.{e.page}] {e.title}: {e.text}" for e in p.narrative_excerpts)
    unit = {"units": "euros, converted to EUR thousands", "thousands": "EUR thousands",
            "millions": "EUR millions, converted to EUR thousands"}
    ext = Extraction(data, p.company_name or name, f"PDF — read by Claude (vision); verify the figures "
                                                   f"({unit[p.unit]})",
                     source_text=f"### Narrative excerpts copied from {name}\n{text}", unmapped=unmapped)
    _fill_review(ext, found_as)
    return ext


# ---------------------------------------------------------------- entry point
def read_financials(file: BinaryIO, name: str, client: anthropic.Anthropic | None = None) -> Extraction:
    """Any supported upload: company financial statements (Excel or PDF) or a completed DD template."""
    if name.lower().endswith(".pdf"):
        return extract_pdf(file.read(), name, client)
    if not name.lower().endswith((".xlsx", ".xlsm")):
        raise ExtractionError("Upload an Excel (.xlsx) or PDF file.")
    wb = load_workbook(file, read_only=True)
    is_template = "Input" in wb.sheetnames
    wb.close()
    file.seek(0)
    if not is_template:
        return extract_excel(file, name)
    data = load_input(file)
    ext = Extraction(data, name, "Completed DD template (Input sheet)")
    _fill_review(ext, {k: "Input sheet" for k in data.lines})
    return ext
