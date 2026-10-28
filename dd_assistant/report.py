"""Layer 4 — write the red flag table (Excel and CSV) and a console summary."""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .metrics import ALL_LABELS, UNITS, Metrics
from .rules import SEVERITIES, RuleResult, alerts, fy, xround
from .units import Money

COLUMNS = ["#", "Rule", "Area", "Years", "Evidence (with source)", "Test and threshold", "What to investigate",
           "Severity", "Question for management", "Document to request", "Reviewer status", "Reviewer comment"]
REVIEW_STATES = ["Pending", "Accepted", "Dismissed"]

# Review of each finding: rule ID -> (status, comment). Not reviewed = ("Pending", "").
Reviews = dict[str, tuple[str, str]]


def fmt_value(v) -> str:
    """Raw value as sent to the AI: None = 'n/a', numbers with up to 4 decimals."""
    if v is None:
        return "n/a"
    if isinstance(v, float):
        return f"{v:.4f}".rstrip("0").rstrip(".")
    return str(v)


def fmt_metric(key: str, v, money: Money | None = None) -> str:
    """Value formatted for people: 23.8%, -3.5 pp, 72 days, 2.65x, 1,415 (in the display unit of money)."""
    return fmt_unit(UNITS.get(key, "eur"), v, money)


def fmt_unit(unit: str, v, money: Money | None = None) -> str:
    if v is None:
        return "n/a"
    if not isinstance(v, (int, float)):
        return str(v)
    if unit == "pct":
        return f"{v * 100:.1f}%"
    if unit == "pp":
        return f"{v:.1f} pp"
    if unit == "days":
        return f"{v:.0f} days"
    if unit == "x":
        return xround(v, 2) + "x"
    if unit == "count":
        return xround(v, 0)
    if money is not None:
        return money.amount(v)
    return f"{round(v):,}".replace("-", "−")


def red_flag_rows(results: list[RuleResult], reviews: Reviews | None = None) -> list[list[str]]:
    reviews = reviews or {}
    rows = []
    for i, r in enumerate(alerts(results), 1):
        rule = r.rule
        state, comment = reviews.get(rule.id, ("Pending", ""))
        rows.append([i, rule.id, rule.area, r.years_text, r.evidence, f"{rule.test}\nRule: {rule.rule_text()}",
                     rule.focus, rule.severity, rule.question, rule.request, state, comment])
    return rows


def summary(results: list[RuleResult]) -> str:
    al = alerts(results)
    by = {k: sum(1 for r in al if r.rule.severity == k) for k in SEVERITIES}
    lines = [f"Findings: {len(al)}  ·  " + "  ·  ".join(f"{k}: {by[k]}" for k in SEVERITIES), ""]
    for i, r in enumerate(al, 1):
        lines.append(f"{i:>2}. {r.rule.id} [{r.rule.severity}] {r.rule.name} ({r.years_text})")
    return "\n".join(lines)


def csv_bytes(results: list[RuleResult], reviews: Reviews | None = None) -> bytes:
    f = io.StringIO(newline="")
    w = csv.writer(f, delimiter=";")
    w.writerow(COLUMNS)
    w.writerows(red_flag_rows(results, reviews))
    return f.getvalue().encode("utf-8-sig")   # utf-8-sig: Excel opens it with the right characters


def write_csv(results: list[RuleResult], path: Path, reviews: Reviews | None = None) -> None:
    Path(path).write_bytes(csv_bytes(results, reviews))


def write_excel(results: list[RuleResult], metrics: Metrics, years: list[int],
                path: Path | BinaryIO, reviews: Reviews | None = None, items: list[dict] | None = None,
                analysis=None) -> None:
    """Write to a file or to an in-memory buffer (the app uses the buffer for downloads).

    items: the extraction review table (what was read and from where), saved as 'Source data'.
    analysis: the whole analysis (pipeline.Analysis): adds the Deal view, Sector and Reading checks sheets.
    """
    money = analysis.money if analysis else Money()
    wb = Workbook()
    ws = wb.active
    ws.title = "Red Flags"
    head = Font(name="Arial", bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3864")
    for c, h in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill = head, fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(red_flag_rows(results, reviews), 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    last = ws.max_row
    dv = DataValidation(type="list", formula1='"' + ",".join(REVIEW_STATES) + '"')
    ws.add_data_validation(dv)
    dv.add(f"K2:K{max(last, 2)}")
    for col, w in zip("ABCDEFGHIJKL", [5, 8, 18, 10, 60, 34, 34, 10, 44, 32, 14, 28]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"

    wm = wb.create_sheet("Metrics")
    wm.append([f"Metric ({money.currency} k unless stated)"] + [fy(y) for y in years])
    for key, label in ALL_LABELS.items():
        wm.append([label] + [("n/a" if metrics[key].get(y) is None else metrics[key].get(y)) for y in years])
    for cell in wm[1]:
        cell.font, cell.fill = head, fill
    wm.column_dimensions["A"].width = 52

    if items:
        ws_src = wb.create_sheet("Source data")
        cols = [c for c in items[0] if c != "key"]
        ws_src.append(cols)
        for it in items:
            ws_src.append([it[c] for c in cols])
        for cell in ws_src[1]:
            cell.font, cell.fill = head, fill
        for col, w in zip("ABCDE", [22, 40, 44, 30, 12]):
            ws_src.column_dimensions[col].width = w
    if analysis:
        _deal_sheet(wb.create_sheet("Deal view"), analysis, head, fill)
        _sector_sheet(wb.create_sheet("Sector"), analysis, head, fill)
        if analysis.checks:
            wc = wb.create_sheet("Reading checks")
            wc.append(["Check", "Year", "Rebuilt from the lines", "As published", "Difference", "OK"])
            for c in analysis.checks:
                wc.append([c["Check"], fy(c["Year"]), c["Rebuilt from the lines"], c["As published"],
                           c["Difference"], "Yes" if c["OK"] else "NO"])
            for cell in wc[1]:
                cell.font, cell.fill = head, fill
            wc.column_dimensions["A"].width = 52
    wb.save(path)


_NUM_FMT = "#,##0;(#,##0)"


def _block(ws, title: str, lines, years: list[int], money: Money, head, fill) -> None:
    """One bridge: label, one column per year (in the display unit), source and note."""
    ws.append([])
    ws.append([title] + [fy(y) for y in years] + ["Source", "Note"])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = head, fill
    for line in lines:
        row = [line.label]
        for y in years:
            x = line.values.get(y)
            row.append(None if x is None else x / 1000 if line.unit == "money" and money.millions
                       else x * 100 if line.unit == "pct" else x)
        ws.append(row + [line.source, line.note])
        bold = line.style in ("subtotal", "total", "header")
        for c, cell in enumerate(ws[ws.max_row]):
            cell.font = Font(name="Arial", size=10, bold=bold)
            if 1 <= c <= len(years) and isinstance(cell.value, (int, float)):
                cell.number_format = _NUM_FMT if line.unit == "money" else "0.0"


def _deal_sheet(ws, a, head, fill) -> None:
    v = a.view
    ws.append([f"Deal view — {a.ext.company} · {a.money.unit_long} unless stated (percentages in %, days, x)"])
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    _block(ws, "Quality of earnings", v.qoe, v.pl_years, a.money, head, fill)
    _block(ws, "Net debt", v.net_debt, v.bs_years, a.money, head, fill)
    _block(ws, "Net working capital", v.nwc, v.bs_years, a.money, head, fill)
    ws.column_dimensions["A"].width = 62
    for col in "BCDE":
        ws.column_dimensions[col].width = 13


def _sector_sheet(ws, a, head, fill) -> None:
    from .sectors import BENCHMARK_SOURCE
    s = a.sector
    ws.append([f"Sector: {s.name} — peers: {s.benchmarks['industry']} ({s.benchmarks['firms']} companies), "
               f"{BENCHMARK_SOURCE}"])
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    from .pipeline import PURPOSES
    c = a.country
    ws.append([f"Purpose: {PURPOSES[a.purpose]} · Country: "
               + (f"{c.name} — {c.framework}; statutory tax rate {c.tax_rate:.2%} ({c.tax_note})"
                  if c and c.tax_rate is not None else c.name if c else "not set")])
    ws.append([])
    ws.append(["Metric", "Company (latest year)", "Sector", "Reading"])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = head, fill
    for b in a.benchmarks:
        ws.append([b["Metric"], fmt_unit(b["unit"], b["Company"]), fmt_unit(b["unit"], b["Sector"]), b["Reading"]])
    ws.append([])
    ws.append(["Calibrated thresholds"])
    ws[ws.max_row][0].font = Font(name="Arial", bold=True)
    for r in a.results:
        if r.rule.calibration:
            ws.append([f"{r.rule.id} · {r.rule.name}", r.rule.limit_text(), "", r.rule.calibration])
    if not any(r.rule.calibration for r in a.results):
        ws.append(["None: the general thresholds apply."])
    ws.append([])
    ws.append(["Sector lens", "Question for management", "", "Document to request"])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = head, fill
    for item in s.lens:
        ws.append([item.area, item.question, "", item.request])
    for col, w in zip("ABCD", [40, 70, 14, 60]):
        ws.column_dimensions[col].width = w
