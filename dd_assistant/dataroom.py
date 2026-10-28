"""Layer 1c — the data room: every document at once, as the seller sends it.

Drop the whole folder:
  - Excel workbooks are read together as one set of financial statements and schedules (P&L in one
    file, balance sheet in another, customers in a third…). Every figure keeps its file and cell.
  - PDFs, Word documents and text files are read as text, on this computer and free: the management
    report, the audit report, contracts. They go to the memo, and the key-document checklist is ticked
    from what is there.
  - An ESEF report package (.zip: the annual report a listed company files in iXBRL) is read from its tags,
    free and exact: see esef.py.
  - Financial statements that exist only as a PDF are read by Claude, on request (paid).
"""
from __future__ import annotations

import html
import io
import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import anthropic
from openpyxl import Workbook, load_workbook

from .esef import EsefReport, is_esef, read_esef
from .extract import Extraction, ExtractionError, _fill_review, _norm, _sheet_type, data_room_text, extract_excel, \
    extract_pdf, read_financials

SPREADSHEETS = (".xlsx", ".xlsm")
DOCUMENTS = (".pdf", ".docx", ".txt", ".md", ".csv")
ESEF = (".zip", ".xhtml", ".html")          # ESEF report package (the report alone is refused with a clear message)

# key documents recognised from a file's name or first lines (Portuguese, Spanish, French, German, Italian, English)
_AGEING = r"aging|ageing|antiguidade|antiguedad|balance agee|altersstruktur|scadenzario|vencimientos"
DOCUMENT_HINTS = {
    "Sales by customer": r"vendas por cliente|ventas por cliente|sales by customer|revenue by customer|ventes par client|"
                         r"umsatz nach kunden|vendite per cliente",
    "Support for EBITDA adjustments (invoices, contracts)": r"fatura|factura|invoice|facture|rechnung|fattur",
    "Breakdown of other income and expenses": r"outros rendimentos|other income|otros ingresos|autres produits|"
                                              r"sonstige betriebliche|altri ricavi",
    "AR ageing": rf"(?:{_AGEING}).*(?:client|receivab|debtor|deudor|creance|forderung|credit)|"
                 rf"(?:client|receivab|debtor|deudor|creance|forderung|credit).*(?:{_AGEING})",
    "Inventory ageing": rf"(?:{_AGEING}).*(?:invent|stock|existenc|lager|vorrat|magazzin|rimanenz)|"
                        rf"(?:invent|stock|existenc|lager|vorrat|magazzin|rimanenz).*(?:{_AGEING})",
    "Debt schedule": r"mapa da divida|mapa de financiamentos|debt schedule|cuadro de deuda|echeancier|schuldenspiegel|"
                     r"piano di ammortamento|loan schedule",
    "Facility agreements": r"contrato de financiamento|facility agreement|loan agreement|credit agreement|"
                           r"contrato de prestamo|contrat de pret|kreditvertrag|contratto di finanziamento",
}


class NeedsPdfReading(ExtractionError):
    """No Excel statements in the data room: the statements must be read from a PDF with Claude (paid)."""

    def __init__(self, pdfs: list[str]):
        self.pdfs = pdfs
        super().__init__("The data room has no Excel financial statements. Read them from a PDF with Claude "
                         f"(uses the Claude API): {', '.join(pdfs)}.")


@dataclass
class _Doc:
    name: str
    text: str
    pages: int


def _pdf_text(content: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    logging.getLogger("pypdf").setLevel(logging.ERROR)   # font warnings are not errors
    reader = PdfReader(io.BytesIO(content))
    pages = [(i, (page.extract_text() or "").strip()) for i, page in enumerate(reader.pages, 1)]
    return "\n".join(f"[p.{i}] {text}" for i, text in pages if text), len(pages)


def _docx_text(content: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
    xml = re.sub(r"</w:p>", "\n", xml)
    return html.unescape(re.sub(r"<[^>]+>", "", xml)).strip()


def read_document(name: str, content: bytes) -> _Doc:
    """The text of a PDF, Word or text file (no AI, no cost)."""
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        text, pages = _pdf_text(content)
        return _Doc(name, text, pages)
    if suffix == ".docx":
        return _Doc(name, _docx_text(content), 0)
    return _Doc(name, content.decode("utf-8", "ignore"), 0)


def _merge(workbooks: list[tuple[str, bytes]]) -> tuple[io.BytesIO, dict[str, str]]:
    """All the sheets of several workbooks in one, cell for cell (so 'DRE!B5' stays 'DRE!B5'), and which file
    each sheet came from."""
    merged, origin = Workbook(), {}
    merged.remove(merged.active)
    for name, content in workbooks:
        wb = load_workbook(io.BytesIO(content), data_only=True)
        for ws in wb.worksheets:
            title, n = ws.title, 2
            while title in origin:
                title, n = f"{ws.title[:27]} ({n})", n + 1
            origin[title] = name
            out = merged.create_sheet(title)
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        out.cell(row=c.row, column=c.column, value=c.value)
    buf = io.BytesIO()
    merged.save(buf)
    buf.seek(0)
    return buf, origin


def _with_files(text: str, origin: dict[str, str]) -> str:
    """'DRE!B5:D5' -> '01 Demonstrações financeiras.xlsx › DRE!B5:D5'."""
    if not text or len(set(origin.values())) < 2:
        return text
    titles = "|".join(re.escape(t) for t in sorted(origin, key=len, reverse=True))
    return re.sub(rf"(?<!\w)({titles})!", lambda m: f"{origin[m.group(1)]} › {m.group(1)}!", text)


def _found_documents(names_and_text: list[tuple[str, str]]) -> set[str]:
    found = set()
    for name, text in names_and_text:
        probe = _norm(f"{Path(name).stem} {text[:1500]}")
        found |= {doc for doc, pattern in DOCUMENT_HINTS.items() if re.search(pattern, probe)}
    return found


def read_data_room(files: list[tuple[str, bytes]], client: anthropic.Anthropic | None = None,
                   read_pdf: str | None = None) -> Extraction:
    """Read every document of a data room. read_pdf: the PDF to read the statements from with Claude, when the
    data room has no Excel statements (it uses the Claude API)."""
    files = sorted(files, key=lambda f: f[0].lower())
    reports = [(n, c) for n, c in files if n.lower().endswith(ESEF) and is_esef(n, c)]
    if len(reports) > 1:
        raise ExtractionError("Upload one ESEF report at a time (each already has two years): "
                              f"{', '.join(n for n, _ in reports)}.")
    esef: EsefReport | None = read_esef(*reports[0]) if reports else None
    sheets = [(n, c) for n, c in files if n.lower().endswith(SPREADSHEETS)]
    docs = [(n, c) for n, c in files if n.lower().endswith(DOCUMENTS)]
    ignored = [n for n, _ in files if not n.lower().endswith(SPREADSHEETS + DOCUMENTS) and
               not (esef and n == esef.name)]
    if not sheets and not docs and not esef:
        raise ExtractionError("Upload Excel (.xlsx), PDF, Word (.docx), text files or an ESEF report package (.zip).")

    templates = [n for n, c in sheets if "Input" in load_workbook(io.BytesIO(c), read_only=True).sheetnames]
    if templates and (len(sheets) > 1 or docs):
        raise ExtractionError(f"{templates[0]} is a completed DD template: upload it on its own.")
    if templates:
        return read_financials(io.BytesIO(sheets[0][1]), sheets[0][0])

    texts = {n: read_document(n, c) for n, c in docs}
    table = []
    if sheets or esef:
        books = sheets + ([(esef.name, esef.book)] if esef else [])
        buf, origin = _merge(books)
        ext = extract_excel(buf, books[0][0])
        prefix = f"{esef.name} › " if esef and len(set(origin.values())) > 1 else ""
        source = lambda s: _with_files(esef.tag(s, prefix) if esef else s, origin)
        d = ext.data
        for key, src in list(d.sources.items()):
            d.sources[key] = source(src)
        d.clients_source, d.adjustments_source = source(d.clients_source), source(d.adjustments_source)
        for b in d.balance_items:
            b.source = source(b.source)
        for row in ext.items + ext.unmapped:
            row["Source"] = source(row["Source"])
        ext.source_text = "\n\n".join(f"## File '{n}'\n{data_room_text(io.BytesIO(c))}" for n, c in sheets)
        if esef:
            ext.company = esef.company
            ext.method = ext.method.replace("Excel — dictionary mapping, figures read from the cells",
                                            "ESEF (iXBRL) — each line read from its tag" + " and Excel" * bool(sheets))
            ext.source_text = f"## ESEF report '{esef.name}'\n{esef.text}\n\n{ext.source_text}".strip()
            table.append({"File": esef.name, "Read as": "ESEF — tagged financial statements (iXBRL)",
                          "Found": f"income statement, balance sheet, cash flow statement ({esef.facts} facts in euros)"})
        for n, _ in sheets:
            kinds = sorted({_KIND_NAMES.get(_sheet_type(ws) or "", "other") for ws in
                            load_workbook(io.BytesIO(dict(sheets)[n]), read_only=True).worksheets})
            table.append({"File": n, "Read as": "Excel — financial data", "Found": ", ".join(kinds)})
    else:
        pdfs = [n for n, _ in docs if n.lower().endswith(".pdf")]
        if not read_pdf:
            raise NeedsPdfReading(pdfs)
        if read_pdf not in dict(docs):
            raise ExtractionError(f"{read_pdf} is not in the data room.")
        ext = extract_pdf(dict(docs)[read_pdf], read_pdf, client)
        table.append({"File": read_pdf, "Read as": "PDF — financial statements read by Claude",
                      "Found": "income statement, balance sheet, cash flow"})

    for n, doc in texts.items():
        if n == read_pdf:
            continue
        ext.source_text += f"\n\n## Document '{n}'\n{doc.text or '(no text: a scanned document?)'}"
        size = (f"{doc.pages} page{'s' * (doc.pages != 1)}" if doc.pages else f"{len(doc.text.split()):,} words")
        table.append({"File": n, "Read as": "Document — text for the memo", "Found": size})
    table += [{"File": n, "Read as": "Not read", "Found": "format not supported"} for n in ignored]

    found = _found_documents([(n, t.text) for n, t in texts.items()] + [(n, "") for n, _ in sheets])
    ext.data.documents = [(doc, ok or doc in found) for doc, ok in ext.data.documents]
    ext.files = table
    if not ext.items:
        _fill_review(ext, {})
    return ext


_KIND_NAMES = {"pl": "income statement", "bs": "balance sheet", "cf": "cash flow statement",
               "customers": "sales by customer", "adjustments": "management adjustments", "debt": "debt schedule",
               "other": "notes"}


def files_in(paths: list[str | Path]) -> list[tuple[str, bytes]]:
    """The files given on the command line: files, or every supported file in a folder."""
    out = []
    for p in map(Path, paths):
        members = sorted(p.iterdir()) if p.is_dir() else [p]
        out += [(f.name, f.read_bytes()) for f in members
                if f.is_file() and f.suffix.lower() in SPREADSHEETS + DOCUMENTS + ESEF and not f.name.startswith("~$")]
    return out
