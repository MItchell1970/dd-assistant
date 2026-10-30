"""Layer 1d — ESEF: the annual financial report every EU listed company files in iXBRL since 2021.

Each figure in an ESEF report is tagged with its IFRS meaning, period, unit and sign, so nothing has
to be guessed: the reader takes the tagged facts of the consolidated primary statements, lays them out
in the order the company presents them, and hands them to the same engine as an Excel data room.
  - IFRS concepts get the name the dictionary knows ('ifrs-full:Revenue' -> 'Revenue');
  - the company's own concepts (extensions) are understood through their ESEF anchoring, the IFRS
    concepts they contain ('Cost of energy sales and other' contains 'RawMaterialsAndConsumablesUsed'),
    or else through their English label;
  - every figure keeps its tag as the source ('ifrs-full:Revenue').
Upload the report package (.zip, from https://filings.xbrl.org or the company's website): it carries the
presentation, labels and anchoring the reader needs. Euros only, like the rest of the engine.
"""
from __future__ import annotations

import io
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import PurePosixPath

from openpyxl import Workbook

from .extract import ExtractionError, _norm

# IFRS concept -> the name the dictionary of extract.py recognises. The order matters when a company's own
# concept contains several of them: it is named after the first ('Purchase of property, plant and equipment
# and purchase of intangible assets').
CANONICAL = {
    # income statement
    "Revenue": "Revenue", "RevenueFromContractsWithCustomers": "Revenue", "CostOfSales": "Cost of sales",
    "RawMaterialsAndConsumablesUsed": "Cost of materials consumed", "CostOfMerchandiseSold": "Cost of goods sold",
    "CostOfPurchasedEnergySold": "Cost of materials consumed",
    "ChangesInInventoriesOfFinishedGoodsAndWorkInProgress": "Change in manufactured inventories",
    "OtherWorkPerformedByEntityAndCapitalised": "Own work capitalised", "GrossProfit": "Gross profit",
    "ServicesExpense": "External supplies and services", "EmployeeBenefitsExpense": "Staff costs",
    "DistributionCosts": "Distribution costs", "AdministrativeExpense": "Administrative expenses",
    "SellingGeneralAndAdministrativeExpense": "Selling, general and administrative",
    "ResearchAndDevelopmentExpense": "Research and development", "OtherIncome": "Other income",
    "OtherOperatingIncomeExpense": "Other operating income", "OtherExpenseByNature": "Other expenses",
    "OtherOperatingExpense": "Other operating expenses",
    "ImpairmentLossImpairmentGainAndReversalOfImpairmentLossDeterminedInAccordanceWithIFRS9":
        "Impairment of receivables",
    "DepreciationAndAmortisationExpense": "Depreciation and amortisation",
    "DepreciationAmortisationAndImpairmentLossReversalOfImpairmentLossRecognisedInProfitOrLoss":
        "Depreciation and amortisation",
    "ShareOfProfitLossOfAssociatesAndJointVenturesAccountedForUsingEquityMethod":
        "Share of profit of associates and joint ventures",
    "ProfitLossFromOperatingActivities": "Operating profit", "FinanceIncome": "Finance income",
    "FinanceCosts": "Finance costs", "ProfitLossBeforeTax": "Profit before tax",
    "IncomeTaxExpenseContinuingOperations": "Income tax", "ProfitLoss": "Profit for the year",
    "ProfitLossAttributableToOwnersOfParent": "Attributable to owners of the parent",
    # balance sheet
    "PropertyPlantAndEquipment": "Property, plant and equipment", "Inventories": "Inventories",
    "CurrentInventoriesHeldForSale": "Inventories", "TradeAndOtherCurrentReceivables": "Trade and other receivables",
    "CurrentTradeReceivables": "Trade receivables", "CashAndCashEquivalents": "Cash and cash equivalents",
    "Equity": "Total equity", "LongtermBorrowings": "Borrowings", "NoncurrentPortionOfNoncurrentBorrowings": "Borrowings",
    "ShorttermBorrowings": "Borrowings", "CurrentBorrowingsAndCurrentPortionOfNoncurrentBorrowings": "Borrowings",
    "CurrentPortionOfLongtermBorrowings": "Borrowings", "TradeAndOtherCurrentPayables": "Trade and other payables",
    "TradeAndOtherCurrentPayablesToTradeSuppliers": "Trade payables", "CurrentTradePayables": "Trade payables",
    "NoncurrentAssets": "Total non-current assets", "CurrentAssets": "Total current assets",
    "NoncurrentLiabilities": "Total non-current liabilities", "CurrentLiabilities": "Total current liabilities",
    "Assets": "Total assets", "EquityAndLiabilities": "Total equity and liabilities",
    "NoncurrentLeaseLiabilities": "Lease liabilities", "CurrentLeaseLiabilities": "Lease liabilities",
    "NoncurrentProvisionsForEmployeeBenefits": "Post-employment benefit obligations",
    "CurrentProvisionsForEmployeeBenefits": "Post-employment benefit obligations",
    "OtherLongtermProvisions": "Provisions", "OtherShorttermProvisions": "Provisions",
    "CurrentTaxLiabilitiesCurrent": "Current tax liabilities", "OtherCurrentFinancialAssets": "Other financial assets",
    "OtherNoncurrentFinancialAssets": "Other financial assets",
    "OtherCurrentFinancialLiabilities": "Other financial liabilities",
    "OtherNoncurrentFinancialLiabilities": "Other financial liabilities",
    "NoncurrentAssetsOrDisposalGroupsClassifiedAsHeldForSale": "Assets held for sale",
    # cash flow statement
    "CashFlowsFromUsedInOperatingActivities": "Net cash from operating activities",
    "CashFlowsFromUsedInOperations": "Cash generated from operations",
    "AdjustmentsForDepreciationAndAmortisationExpense": "Depreciation and amortisation",
    "AdjustmentsForDepreciationAndAmortisationExpenseAndImpairmentLossReversalOfImpairmentLossRecognisedInProfitOrLoss":
        "Depreciation and amortisation",
    "PurchaseOfPropertyPlantAndEquipmentClassifiedAsInvestingActivities": "Purchase of property, plant and equipment",
    "PurchaseOfIntangibleAssetsClassifiedAsInvestingActivities": "Purchase of intangible assets",
    "ProceedsFromSalesOfPropertyPlantAndEquipmentClassifiedAsInvestingActivities":
        "Proceeds from sale of property, plant and equipment",
}
# lines a company's own concept may be named after when it contains them (not totals, not items to assess)
CORE = [k for k, v in CANONICAL.items() if v not in (
    "Lease liabilities", "Post-employment benefit obligations", "Provisions", "Current tax liabilities",
    "Other financial assets", "Other financial liabilities", "Assets held for sale", "Total equity",
    "Total non-current assets", "Total current assets", "Total non-current liabilities", "Total current liabilities",
    "Total assets", "Total equity and liabilities")]
# section headers of the balance sheet, from the IFRS abstract concepts (they carry no figure)
HEADERS = {"NoncurrentAssetsAbstract": "Non-current assets", "CurrentAssetsAbstract": "Current assets",
           "EquityAbstract": "Equity", "EquityAndLiabilitiesAbstract": "Equity and liabilities",
           "LiabilitiesAbstract": "Liabilities", "NoncurrentLiabilitiesAbstract": "Non-current liabilities",
           "CurrentLiabilitiesAbstract": "Current liabilities"}
TITLES = {"pl": "ESEF income statement", "bs": "ESEF balance sheet", "cf": "ESEF cash flow"}
LABEL_ROLES = ("/label", "/terseLabel", "/totalLabel", "/verboseLabel", "/periodEndLabel", "/negatedLabel",
               "/negatedTerseLabel", "/documentation")
ENTITY = {"NameOfReportingEntityOrOtherMeansOfIdentification": "Name of the reporting entity",
          "LegalFormOfEntity": "Legal form", "CountryOfIncorporation": "Country of incorporation",
          "DomicileOfEntity": "Domicile", "PrincipalPlaceOfBusiness": "Principal place of business",
          "DescriptionOfNatureOfEntitysOperationsAndPrincipalActivities": "Nature of business"}
_ATTRS = re.compile(r"([\w:.-]+)=\"([^\"]*)\"")


@dataclass
class EsefReport:
    name: str
    company: str
    book: bytes                          # the primary statements as a workbook for the engine
    where: dict[tuple[str, int], str]    # (sheet, row) -> the tag of that line
    facts: int                           # tagged facts in euros used
    text: str                            # statements and entity information, for the memo

    def tag(self, source: str, prefix: str = "") -> str:
        """'ESEF income statement!B7:C7' -> 'ifrs-full:Revenue': the tag is the source."""
        titles = "|".join(map(re.escape, TITLES.values()))

        def sub(m):
            concept = self.where.get((m.group(1), int(m.group(2))))
            return prefix + concept if concept else m.group(0)
        return re.sub(rf"(?<!\w)({titles})!B(\d+):[A-Z]+\d+", sub, source)


def _attrs(tag: str) -> dict[str, str]:
    return dict(_ATTRS.findall(tag))


def _local(href: str) -> str:
    """'ifrs-full_Revenue' or 'ifrs-full:Revenue' -> 'Revenue'."""
    return re.split(r"[_:]", href, maxsplit=1)[-1]


def _qname(href: str) -> str:
    return href.replace("_", ":", 1)


def _humanise(local: str) -> str:
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", local).capitalize()


def _clean(label: str) -> str:
    return re.sub(r"\s*\[(?:abstract|line items|table)\]\s*$", "", label, flags=re.I).strip()


def is_esef(name: str, content: bytes) -> bool:
    """An ESEF report package (.zip with reports/*.xhtml) or an inline XBRL report on its own."""
    low = name.lower()
    if low.endswith(".zip"):
        try:
            return any(_is_report(n) for n in zipfile.ZipFile(io.BytesIO(content)).namelist())
        except zipfile.BadZipFile:
            return False
    return low.endswith((".xhtml", ".html", ".htm")) and b"<ix:header" in content


def _is_report(member: str) -> bool:
    return member.lower().endswith((".xhtml", ".html")) and "/reports/" in f"/{member.lower()}"


def _package(name: str, content: bytes) -> tuple[str, dict[str, str]]:
    """The report and the company's taxonomy files (labels, presentation, anchoring)."""
    z = zipfile.ZipFile(io.BytesIO(content))
    report = next(n for n in z.namelist() if _is_report(n))
    taxonomy = {PurePosixPath(n).name: z.read(n).decode("utf-8", "ignore") for n in z.namelist()
                if re.search(r"_(lab-[\w-]+|pre|def)\.xml$", n.lower())}
    return z.read(report).decode("utf-8", "ignore"), taxonomy


def _periods(text: str) -> tuple[dict[str, tuple[str, int]], dict[str, str]]:
    """Contexts without dimensions (the consolidated figures) at the financial year end: id -> ('instant' or
    'duration', fiscal year); and units: id -> measure."""
    raw = []
    for m in re.finditer(r"<xbrli:context\b([^>]*)>(.*?)</xbrli:context>", text, re.S):
        body = m.group(2)
        if "segment" in body or "scenario" in body:
            continue
        cid = _attrs(m.group(1)).get("id")
        if (inst := re.search(r"<xbrli:instant>\s*(\d{4}-\d\d-\d\d)", body)):
            raw.append((cid, "instant", None, date.fromisoformat(inst.group(1))))
        elif (start := re.search(r"<xbrli:startDate>\s*(\d{4}-\d\d-\d\d)", body)) and \
                (end := re.search(r"<xbrli:endDate>\s*(\d{4}-\d\d-\d\d)", body)):
            raw.append((cid, "duration", date.fromisoformat(start.group(1)), date.fromisoformat(end.group(1))))
    years = [e for _, kind, s, e in raw if kind == "duration" and 350 <= (e - s).days <= 380]
    if not years:
        raise ExtractionError("No full financial year found in this ESEF report.")
    year_end = max(years)

    def at_year_end(d: date) -> bool:                  # 31 December, or the company's own year end (52/53 weeks)
        try:
            return abs((d.replace(year=year_end.year) - year_end).days) <= 7
        except ValueError:                             # 29 February
            return False

    contexts = {cid: (kind, e.year) for cid, kind, s, e in raw if at_year_end(e) and
                (kind == "instant" or 350 <= (e - s).days <= 380)}
    units = {}
    for m in re.finditer(r"<xbrli:unit\b([^>]*)>(.*?)</xbrli:unit>", text, re.S):
        units[_attrs(m.group(1)).get("id")] = "/".join(re.findall(r"<xbrli:measure>\s*([^<\s]+)", m.group(2)))
    return contexts, units


def _value(raw: str, a: dict[str, str]) -> float | None:
    """The fact's value in EUR thousands: displayed number, transformation format, scale and sign."""
    fmt = a.get("format", "").lower()
    shown = re.sub(r"<[^>]+>", "", raw).strip()
    if "zero" in fmt or shown in ("-", "—", "–", ""):  # ixt:fixed-zero, ixt:zerodash
        number = 0.0
    else:
        digits = re.sub(r"[^\d.,]", "", shown)
        if "comma" in fmt:                             # num-comma-decimal: 1.234.567,8
            digits = digits.replace(".", "").replace(",", ".")
        else:                                          # num-dot-decimal: 1,234,567.8
            digits = digits.replace(",", "")
        try:
            number = float(digits)
        except ValueError:
            return None
    number *= 10 ** int(a.get("scale") or 0) / 1000
    return -number if a.get("sign") == "-" else number


def _facts(text: str, contexts: dict, units: dict) -> tuple[dict[str, dict[str, dict[int, float]]], int, dict]:
    """Monetary facts in euros: concept -> period type -> {year: EUR k}; how many; and the entity's text facts."""
    best: dict[tuple[str, str, int], tuple[float, float]] = {}
    currencies: dict[str, int] = defaultdict(int)
    for m in re.finditer(r"<ix:nonFraction\b([^>]*)>(.*?)</ix:nonFraction>", text, re.S):
        a = _attrs(m.group(1))
        unit = units.get(a.get("unitRef"), "")
        if a.get("contextRef") not in contexts or not re.fullmatch(r"iso4217:[A-Z]{3}", unit):
            continue                                   # segments, shares, ratios, per-share amounts
        currencies[unit[-3:]] += 1
        v = _value(m.group(2), a)
        if unit != "iso4217:EUR" or v is None:
            continue
        kind, year = contexts[a["contextRef"]]
        precision = 99.0 if a.get("decimals", "").upper() == "INF" else float(a.get("decimals") or -99)
        key = (a["name"], kind, year)
        if key not in best or precision > best[key][0]:  # the same figure rounded in the management report
            best[key] = (precision, v)
    if currencies and max(currencies, key=currencies.get) != "EUR":
        raise ExtractionError(f"This ESEF report is in {max(currencies, key=currencies.get)}. DD Assistant analyses "
                              "statements in euros (European market).")
    values: dict[str, dict[str, dict[int, float]]] = defaultdict(lambda: defaultdict(dict))
    for (concept, kind, year), (_, v) in best.items():
        values[concept][kind][year] = v
    entity = {}
    for m in re.finditer(r"<ix:nonNumeric\b([^>]*)>(.*?)</ix:nonNumeric>", text, re.S):
        local = _local(_attrs(m.group(1)).get("name", ""))
        if local in ENTITY and local not in entity:
            entity[local] = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", m.group(2))).strip()[:400]
    if (lei := re.search(r"<xbrli:identifier\b[^>]*>\s*([A-Z0-9]{20})\s*<", text)):
        entity["LEI"] = lei.group(1)
    return values, len(best), entity


def _labels(taxonomy: dict[str, str]) -> dict[str, str]:
    """The company's labels for its own concepts: English first, then any other language."""
    out: dict[str, str] = {}
    for name in sorted((n for n in taxonomy if "_lab-" in n), key=lambda n: ("_lab-en" not in n, n)):
        lab = taxonomy[name]
        locs = {_attrs(t)["xlink:label"]: _attrs(t)["xlink:href"].split("#")[-1]
                for t in re.findall(r"<link:loc\b[^>]*>", lab)}
        texts: dict[str, dict[str, str]] = defaultdict(dict)
        for m in re.finditer(r"<link:label\b([^>]*)>(.*?)</link:label>", lab, re.S):
            a = _attrs(m.group(1))
            role = next((r for r in LABEL_ROLES if a.get("xlink:role", "").endswith(r)), None)
            if role:
                texts[a["xlink:label"]][role] = re.sub(r"\s+", " ", m.group(2)).strip()
        for arc in re.findall(r"<link:labelArc\b[^>]*>", lab):
            a = _attrs(arc)
            concept, found = locs.get(a.get("xlink:from"), ""), texts.get(a.get("xlink:to"), {})
            if concept not in out and (best := next((found[r] for r in LABEL_ROLES if r in found), None)):
                out[concept] = best
    return out


def _anchors(taxonomy: dict[str, str]) -> dict[str, list[str]]:
    """The company's own concept -> the IFRS concepts it contains (ESEF wider-narrower anchoring)."""
    out: dict[str, list[str]] = defaultdict(list)
    for name, d in taxonomy.items():
        if not name.endswith("_def.xml"):
            continue
        for link in re.finditer(r"<link:definitionLink\b[^>]*>(.*?)</link:definitionLink>", d, re.S):
            locs = {_attrs(t)["xlink:label"]: _attrs(t)["xlink:href"].split("#")[-1]
                    for t in re.findall(r"<link:loc\b[^>]*>", link.group(1))}
            for t in re.findall(r"<link:definitionArc\b[^>]*>", link.group(1)):
                a = _attrs(t)
                if a.get("xlink:arcrole", "").endswith("wider-narrower"):
                    out[locs[a["xlink:from"]]].append(locs[a["xlink:to"]])
    return out


def _kind(root: str) -> tuple[str, int] | None:
    """Which primary statement a presentation tree is, from its root ('StatementOfCashFlowsAbstract'), and a
    preference: the income statement before a statement of comprehensive income."""
    r = _local(root).lower()
    if not r.startswith(("statementof", "incomestatement", "balancesheet")):
        return None                                    # notes ('DisclosureOf…'), changes in equity
    if "cashflow" in r:
        return "cf", 0
    if "financialposition" in r or "balancesheet" in r:
        return "bs", 0
    if "incomestatement" in r or "profitorloss" in r:
        return "pl", 0
    return ("pl", 1) if "comprehensiveincome" in r else None


def _statements(taxonomy: dict[str, str]) -> dict[str, list[tuple[int, str]]]:
    """Each primary statement as (depth, concept) in presentation order, from the presentation linkbase."""
    pre = next((t for n, t in taxonomy.items() if n.endswith("_pre.xml")), "")
    found: dict[str, tuple[int, list[tuple[int, str]]]] = {}
    for role in re.finditer(r"<link:presentationLink\b[^>]*>(.*?)</link:presentationLink>", pre, re.S):
        body = role.group(1)
        locs = {_attrs(t)["xlink:label"]: _attrs(t)["xlink:href"].split("#")[-1]
                for t in re.findall(r"<link:loc\b[^>]*>", body)}
        kids, children = defaultdict(list), set()
        for t in re.findall(r"<link:presentationArc\b[^>]*>", body):
            a = _attrs(t)
            kids[a["xlink:from"]].append((float(a.get("order") or 0), a["xlink:to"]))
            children.add(a["xlink:to"])
        roots = [n for n in kids if n not in children]
        kind = next((k for r in roots if (k := _kind(locs[r]))), None)
        if not kind or (kind[0] in found and found[kind[0]][0] <= kind[1]):
            continue                                   # the first statement of each kind: the consolidated one
        rows: list[tuple[int, str]] = []

        def walk(node, depth):
            rows.append((depth, locs[node]))
            for _, k in sorted(kids[node]):
                walk(k, depth + 1)

        for r in roots:
            walk(r, 0)
        found[kind[0]] = (kind[1], rows)
    return {k: rows for k, (_, rows) in found.items()}


def read_esef(name: str, content: bytes) -> EsefReport:
    """The consolidated primary statements of an ESEF report package, ready for the engine."""
    from .report import fmt_value

    if not name.lower().endswith(".zip"):
        raise ExtractionError(f"{name} is the ESEF report on its own. Upload the report package (.zip) instead — "
                              "on filings.xbrl.org it is the 'package' link: it carries the order of the statements "
                              "and the company's own labels and anchoring.")
    text, taxonomy = _package(name, content)
    contexts, units = _periods(text)
    values, count, entity = _facts(text, contexts, units)
    statements = _statements(taxonomy)
    if "pl" not in statements or "bs" not in statements:
        raise ExtractionError(f"No income statement or balance sheet found in the ESEF report {name}.")
    labels, anchors = _labels(taxonomy), _anchors(taxonomy)

    def label(href: str) -> str:
        local = _local(href)
        if href.startswith("ifrs-full") and local in CANONICAL:
            return CANONICAL[local]
        # a company concept that contains IFRS lines is named after them
        inside = sorted({_local(n) for n in anchors.get(href, []) if _local(n) in CORE}, key=CORE.index)
        names = list(dict.fromkeys(CANONICAL[c] for c in inside))
        own = _clean(labels.get(href) or _humanise(local))
        if names:                                      # ... and the company's own name stays visible
            text = names[0] + "".join(f" and {n[0].lower()}{n[1:]}" for n in names[1:])
            return text if _norm(own) == _norm(text) else f"{text} ({own})"
        return own

    wb, where = Workbook(), {}
    ws = wb.active
    ws.title = "ESEF entity"
    company = entity.get("NameOfReportingEntityOrOtherMeansOfIdentification") or PurePosixPath(name).stem
    ws.append([company])
    info = [(ENTITY.get(k, k), v) for k, v in entity.items()]
    for key, value in info:
        if len(value) <= 80:                            # the description of the business goes to the memo only
            ws.append([key, value])
    out = [f"ESEF annual financial report (iXBRL) of {company}. Figures in EUR thousands, read from the tags; "
           "the tag of each line is in brackets."] + [f"{k} | {v}" for k, v in info]
    for kind in ("pl", "bs", "cf"):
        period = "instant" if kind == "bs" else "duration"
        anchor = {"pl": "Revenue", "bs": "Assets", "cf": "CashFlowsFromUsedInOperatingActivities"}[kind]
        years = sorted(values.get(f"ifrs-full:{anchor}", {}).get(period, {})) or \
            sorted({y for c in values.values() for y in c.get(period, {})})
        ws = wb.create_sheet(TITLES[kind])
        ws.append(["EUR thousands"] + years)
        out += ["", f"### {TITLES[kind].removeprefix('ESEF ').capitalize()} (EUR thousands)"]

        def line(text_label, figures, tag):
            ws.append([text_label] + figures)
            where[(ws.title, ws.max_row)] = tag
            out.append(f"{text_label} [{tag}] | " + " | ".join(f"FY{str(y)[2:]}: {fmt_value(v)}"
                                                               for y, v in zip(years, figures) if v is not None))

        for depth, href in statements.get(kind, []):
            local, facts = _local(href), values.get(_qname(href), {}).get(period, {})
            if local.endswith("Abstract"):
                if kind != "pl":                        # sections of the balance sheet and the cash flow
                    ws.append([HEADERS.get(local) or _clean(labels.get(href) or _humanise(local[:-8]))])
                    out.append(ws.cell(ws.max_row, 1).value + ":")
                continue
            if any(y in facts for y in years) and not (kind == "cf" and local == "CashAndCashEquivalents"):
                line(label(href), [facts.get(y) for y in years], _qname(href))
        if kind == "cf":                                # opening and closing cash, from the balance sheet dates
            cash = values.get("ifrs-full:CashAndCashEquivalents", {}).get("instant", {})
            line("Cash and cash equivalents at beginning of period", [cash.get(y - 1) for y in years],
                 "ifrs-full:CashAndCashEquivalents (opening)")
            line("Cash and cash equivalents at end of period", [cash.get(y) for y in years],
                 "ifrs-full:CashAndCashEquivalents")
    buf = io.BytesIO()
    wb.save(buf)
    return EsefReport(name, company, buf.getvalue(), where, count, "\n".join(out))
