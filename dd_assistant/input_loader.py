"""Layer 1 — company data: the standard model and the reader for the Excel DD template.

CompanyData is the standard model every later layer works on. It is filled either from
the Input sheet of the DD template (this module) or straight from the company's own
financial statements (extract.py). In the template, rows are found by the text in
column A (not by row number), so the template can gain rows without breaking the code.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

# Column A text in the template's Input sheet -> key used in the code.
LINE_LABELS = {
    "Vendas e serviços prestados": "vendas",
    "Custo das mercadorias vendidas (CMV)": "cmv",
    "Fornecimentos e serviços externos": "fse",
    "Gastos com o pessoal": "pess",
    "Imparidade de dívidas a receber": "imp",
    "Outros rendimentos": "orend",
    "Outros gastos": "ogast",
    "Depreciações e amortizações": "da",
    "Juros e gastos similares": "juros",
    "Imposto sobre o rendimento": "irc",
    "Resultado líquido (DRE)": "rl_dre",
    "Ativos fixos tangíveis": "aft",
    "Outros ativos não correntes": "oanc",
    "Inventários": "inv",
    "Clientes": "cli",
    "Outros ativos correntes": "oac",
    "Caixa e depósitos bancários": "caixa",
    "Total do capital próprio": "cp",
    "Resultado líquido (balanço)": "rl_bs",
    "Financiamentos obtidos (não corrente)": "fin_lp",
    "Outros passivos não correntes": "opnc",
    "Fornecedores": "forn",
    "Financiamentos obtidos (corrente)": "fin_cp",
    "Outros passivos correntes": "opc",
    "Resultado líquido (DFC)": "rl_dfc",
    "Variação de clientes": "d_cli",
    "Variação de inventários": "d_inv",
    "Variação de fornecedores": "d_forn",
    "Variação de outros ativos correntes": "d_oac",
    "Variação de outros passivos correntes": "d_opc",
    "Caixa gerada pelas operações (antes de juros e impostos)": "cgo",
    "Fluxo de caixa das atividades operacionais (FCO)": "fco",
    "Pagamentos de ativos fixos (capex)": "capex",
    "Recebimentos da alienação de ativos fixos": "alien",
    "(Ganhos)/perdas na alienação de ativos fixos": "ganho_alien",
    "Caixa no início do período": "cx_ini",
    "Caixa no fim do período": "cx_fim",
    "Total em dívida segundo o mapa": "map_tot",
    "Parcela corrente segundo o mapa": "map_cur",
    "Covenant: dívida líquida / EBITDA máximo": "cov_lim",
}

# Start of the section headers (the "(" avoids confusion with document-list rows
# such as "Vendas por cliente").
SECTION_CLIENTS = "Vendas por cliente ("
SECTION_ADJ = "Ajustes ao EBITDA apresentados"
SECTION_DOCS = "Documentos recebidos ("
OTHER_CLIENTS = "Outros clientes"

# The ten key documents the agent expects in a financial data room (R18).
KEY_DOCUMENTS = [
    "Income statement", "Balance sheet", "Cash flow statement", "Sales by customer",
    "Support for EBITDA adjustments (invoices, contracts)", "Breakdown of other income and expenses",
    "AR ageing", "Inventory ageing", "Debt schedule", "Facility agreements",
]
# Document names in the Portuguese template -> English.
DOC_NAMES = dict(zip([
    "Demonstração de resultados", "Balanço", "Demonstração de fluxos de caixa", "Vendas por cliente",
    "Evidência dos ajustes ao EBITDA (faturas, contratos)", "Detalhe de outros rendimentos e outros gastos",
    "Aging de clientes", "Aging de inventários", "Mapa da dívida", "Contratos de financiamento",
], KEY_DOCUMENTS))

YEAR_COLS = ["B", "C", "D", "E"]  # year columns in the Input sheet
HEADER_ROW = 4
SOURCE_COL = "F"


@dataclass
class Named:
    """A named line (customer or adjustment) with one value per year."""
    name: str
    values: dict[int, float | None]


@dataclass
class BalanceItem:
    """A balance sheet line that may belong in net debt: debt-like (leases, pensions, tax payable…) or
    cash-like (marketable securities…). Found in the statements; the reviewer decides whether to include it."""
    category: str                      # e.g. 'Lease liabilities'
    label: str                         # as written in the file
    nature: str                        # 'debt' or 'cash'
    kind: str                          # 'current', 'noncurrent' or ''
    values: dict[int, float | None]
    source: str
    include: bool                      # default: a prudent buy-side view
    in_nwc: bool = False               # sits inside other current assets/liabilities (excluded from NWC if included)


@dataclass
class CompanyData:
    years: list[int]
    lines: dict[str, dict[int, float | None]] = field(default_factory=dict)
    sources: dict[str, str] = field(default_factory=dict)
    clients: list[Named] = field(default_factory=list)
    other_clients: dict[int, float | None] = field(default_factory=dict)
    clients_source: str = ""
    adjustments: list[Named] = field(default_factory=list)
    adjustments_source: str = ""
    documents: list[tuple[str, bool]] = field(default_factory=list)
    currency: str = "EUR"              # amounts in thousands of this currency
    pl_format: str = "nature"          # 'nature' (costs by type, as SNC) or 'function' (cost of sales, SG&A…)
    balance_items: list[BalanceItem] = field(default_factory=list)

    def v(self, key: str, year: int) -> float | None:
        """Value of a line in a year (None if missing)."""
        return self.lines.get(key, {}).get(year)

    @property
    def pl_years(self) -> list[int]:
        """Years with an income statement (revenue present)."""
        return [y for y in self.years if self.v("vendas", y) is not None]

    @property
    def bs_years(self) -> list[int]:
        """Years with a balance sheet (receivables, equity or cash present)."""
        return [y for y in self.years if any(self.v(k, y) is not None for k in ("cli", "cp", "caixa"))]


def _num(x):
    return float(x) if isinstance(x, (int, float)) else None


def load_input(path: str | Path | BinaryIO) -> CompanyData:
    """Read the template's Input sheet (file or in-memory buffer) and return the company data."""
    wb = load_workbook(path, data_only=True)
    ws = wb["Input"]
    years = [int(ws[f"{c}{HEADER_ROW}"].value) for c in YEAR_COLS]
    year_of = dict(zip(YEAR_COLS, years))
    data = CompanyData(years=years)

    section = None
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        label = ws[f"A{row}"].value
        label = label.strip() if isinstance(label, str) else label
        values = {year_of[c]: _num(ws[f"{c}{row}"].value) for c in YEAR_COLS}
        source = ws[f"{SOURCE_COL}{row}"].value or ""

        # section headers switch the reading mode
        if isinstance(label, str) and label.startswith(SECTION_CLIENTS):
            section = "clients"; continue
        if isinstance(label, str) and label.startswith(SECTION_ADJ):
            section = "adj"; continue
        if isinstance(label, str) and label.startswith(SECTION_DOCS):
            section = "docs"; continue

        if section == "clients":
            if label == OTHER_CLIENTS:
                data.other_clients = values
                section = None
                continue
            if not data.clients_source:
                data.clients_source = source
            if label:  # empty template rows are skipped
                data.clients.append(Named(str(label), values))
            continue
        if section == "adj":
            if not data.adjustments_source and source:
                data.adjustments_source = source
            if label:
                data.adjustments.append(Named(str(label), values))
            elif source == "" and all(v is None for v in values.values()):
                section = None  # a fully empty row closes the section
            continue
        if section == "docs":
            if label:
                data.documents.append((DOC_NAMES.get(str(label), str(label)), str(ws[f"B{row}"].value).strip() == "Sim"))
            continue

        if label in LINE_LABELS:
            key = LINE_LABELS[label]
            data.lines[key] = values
            data.sources[key] = source
    return data
