"""Extraction from the company's own files: reading the raw data room must give exactly the
same figures, customers, adjustments and findings as the template completed by hand."""
import io
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from dd_assistant.extract import _from_pdf, _PdfStatements, ExtractionError, apply_edits, extract_excel, read_financials
from dd_assistant.input_loader import load_input
from dd_assistant.metrics import compute_metrics
from dd_assistant.rules import alerts, evaluate

DATA = Path(__file__).resolve().parents[1] / "data"
PAIRS = [("Duriva_data_room.xlsx", "DD_template_modulo1.xlsx"), ("Velmora_data_room.xlsx", "DD_template_velmora.xlsx")]


def _filled(values: dict) -> dict:
    return {y: v for y, v in values.items() if v is not None}


class RawDataRoomTest(unittest.TestCase):
    def test_same_figures_as_the_hand_filled_template(self):
        for room, template in PAIRS:
            a, b = extract_excel(DATA / room).data, load_input(DATA / template)
            with self.subTest(company=room):
                self.assertEqual(a.years, b.years)
                for key in b.lines:
                    self.assertEqual(_filled(a.lines.get(key, {})), _filled(b.lines[key]), key)
                named = lambda xs: [(x.name, _filled(x.values)) for x in xs]
                self.assertEqual(named(a.clients), named(b.clients))
                self.assertEqual(named(a.adjustments), named(b.adjustments))
                self.assertEqual(_filled(a.other_clients), _filled(b.other_clients))

    def test_same_findings_as_the_template_for_duriva(self):
        a = extract_excel(DATA / "Duriva_data_room.xlsx").data
        b = load_input(DATA / "DD_template_modulo1.xlsx")
        ids = lambda d: [r.rule.id for r in alerts(evaluate(d, compute_metrics(d)))]
        self.assertEqual(ids(a), ids(b))
        self.assertEqual(len(ids(a)), 15)

    def test_every_line_has_a_cell_source(self):
        ext = extract_excel(DATA / "Duriva_data_room.xlsx")
        self.assertEqual(ext.missing, [])
        self.assertEqual(ext.data.sources["vendas"], "DRE!B5:D5")
        self.assertEqual(ext.data.sources["cov_lim"], "Dívida!I6")
        self.assertEqual(ext.company, "Duriva Distribuição Industrial, S.A.")
        found = {i["key"]: i["Found as"] for i in ext.items}
        self.assertEqual(found["vendas"], "Vendas e serviços prestados")

    def test_documents_not_in_the_data_room_are_flagged(self):
        docs = dict(extract_excel(DATA / "Velmora_data_room.xlsx").data.documents)
        self.assertTrue(docs["Income statement"])
        self.assertFalse(docs["AR ageing"])                     # Velmora's data room has no ageing

    def test_english_labels_and_positive_costs(self):
        """An English P&L with costs shown as positive numbers is mapped and the costs become negative."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Income statement"
        ws.append(["EUR k", "FY2024", "FY2025"])
        for row in [["Revenue", 1000, 1100], ["Cost of sales", 600, 650], ["Staff costs", 200, 210],
                    ["Profit for the year", 150, 170]]:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        d = extract_excel(buf).data
        self.assertEqual(d.lines["vendas"], {2024: 1000.0, 2025: 1100.0})
        self.assertEqual(d.lines["cmv"], {2024: -600.0, 2025: -650.0})
        self.assertEqual(d.lines["pess"], {2024: -200.0, 2025: -210.0})
        self.assertEqual(d.lines["rl_dre"], {2024: 150.0, 2025: 170.0})

    def test_not_a_financial_file(self):
        wb = Workbook()
        wb.active.append(["hello", "world"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        with self.assertRaises(ExtractionError):
            extract_excel(buf)

    def test_template_upload_still_works(self):
        with open(DATA / "DD_template_modulo1.xlsx", "rb") as fh:
            ext = read_financials(fh, "DD_template_modulo1.xlsx")
        self.assertIn("template", ext.method)

    def test_reviewer_edits_change_the_figures(self):
        d = extract_excel(DATA / "Duriva_data_room.xlsx").data
        edited = apply_edits(d, {"orend": {2025: 60.0}})
        self.assertEqual(edited.v("orend", 2025), 60.0)
        self.assertEqual(d.v("orend", 2025), 310.0)              # the original is untouched


def _workbook(sheets: dict[str, list[list]]) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _ids(d):
    return [r.rule.id for r in alerts(evaluate(d, compute_metrics(d)))]


class PublicAccountsTest(unittest.TestCase):
    """Published accounts have no sales by customer, no management adjustments and no debt schedule."""

    def test_data_room_without_customers_adjustments_or_debt_schedule(self):
        wb = load_workbook(DATA / "Duriva_data_room.xlsx")
        for name in ("Vendas por Cliente", "EBITDA Ajustado", "Dívida"):
            del wb[name]
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        d = extract_excel(buf).data
        m = compute_metrics(d)
        self.assertIsNone(m["top_share"].get(2025))            # n/a, not an error
        self.assertIsNone(m["ebitda_adj"].get(2025))           # no adjusted EBITDA presented
        ids = _ids(d)
        for rule in ("R01", "R02", "R05", "R06", "R16", "R17", "R20"):   # need the missing schedules
            self.assertNotIn(rule, ids)
        self.assertIn("R18", ids)                               # the missing documents are a finding
        self.assertIn("R10", ids)                               # the statements' own findings remain

    def test_portuguese_annual_report_layout(self):
        """Two years, 'Financiamentos obtidos' in both liability sections, direct-method cash flow."""
        buf = _workbook({
            "DRE": [["Rubrica", 2025, 2024],
                    ["Vendas e prestações de serviços", 11000, 10000],
                    ["Custo das vendas", -6600, -6000],
                    ["Fornecimentos e serviços externos", -1600, -1500],
                    ["Gastos com o pessoal", -1300, -1200],
                    ["Gastos/reversões de depreciação e de amortização", -320, -300],
                    ["Gastos financeiros", -90, -100],
                    ["Imposto sobre o rendimento do exercício", -230, -200],
                    ["Resultado líquido do exercício", 860, 700]],
            "Balanço": [["Rubrica", 2025, 2024],
                        ["Ativo não corrente"],
                        ["Ativos fixos tangíveis", 3100, 3000],
                        ["Ativo corrente"],
                        ["Inventários", 1600, 1500],
                        ["Clientes", 1950, 1800],
                        ["Caixa e depósitos bancários", 600, 500],
                        ["Capital próprio"],
                        ["Resultado líquido do exercício", 860, 700],
                        ["Total do capital próprio", 4150, 3600],
                        ["Passivo não corrente"],
                        ["Financiamentos obtidos", 1600, 1800],
                        ["Passivo corrente"],
                        ["Fornecedores", 1100, 1000],
                        ["Financiamentos obtidos", 400, 400]],
            "DFC": [["Rubrica", 2025, 2024],
                    ["Atividades operacionais"],
                    ["Recebimentos de clientes", 10850, 9800],
                    ["Pagamentos a fornecedores", -8100, -7300],
                    ["Pagamentos ao pessoal", -1300, -1200],
                    ["Caixa gerada pelas operações", 1450, 1300],
                    ["Fluxos de caixa das atividades operacionais (1)", 1220, 1100],
                    ["Pagamentos respeitantes a:"],
                    ["Ativos fixos tangíveis", -420, -400],
                    ["Recebimentos provenientes de:"],
                    ["Ativos fixos tangíveis", 0, 0],
                    ["Caixa e seus equivalentes no início do período", 500, 450],
                    ["Caixa e seus equivalentes no fim do período", 600, 500]],
        })
        ext = extract_excel(buf)
        d = ext.data
        self.assertEqual(d.lines["fin_lp"], {2024: 1800.0, 2025: 1600.0})   # read with the section header
        self.assertEqual(d.lines["fin_cp"], {2024: 400.0, 2025: 400.0})
        self.assertEqual(d.lines["capex"], {2024: -400.0, 2025: -420.0})
        self.assertEqual(d.lines["alien"], {2024: 0.0, 2025: 0.0})
        self.assertEqual(d.lines["juros"], {2024: -100.0, 2025: -90.0})
        self.assertEqual(d.lines["cx_fim"], {2024: 500.0, 2025: 600.0})
        m = compute_metrics(d)
        self.assertAlmostEqual(m["conv"][2025], 1450 / 1500)
        self.assertIsNone(m["wc_diff"].get(2025))              # no working capital lines to compare
        self.assertEqual(m["ni_diff"][2025], 0)                # compared with the balance sheet only
        self.assertEqual(m["net_debt"][2025], 1400)
        self.assertEqual(_ids(d), ["R18"])                     # no false alarms on a clean set of accounts


class PdfConversionTest(unittest.TestCase):
    """What Claude returns for a PDF is converted by the code: units and signs, never by the AI."""

    def test_millions_become_thousands_and_costs_negative(self):
        yv = lambda *pairs: [{"year": y, "value": v} for y, v in pairs]
        parsed = _PdfStatements.model_validate({
            "company_name": "Acme", "unit": "millions",
            "line_items": [
                {"key": "vendas", "label_in_document": "Revenue", "page": 3, "values": yv((2024, 10.0), (2025, 12.5))},
                {"key": "cmv", "label_in_document": "Cost of sales", "page": 3, "values": yv((2024, 6.0), (2025, 7.0))},
                {"key": "cli", "label_in_document": "Trade receivables", "page": 4, "values": yv((2024, 2.0), (2025, 2.5))},
            ],
            "customers": [], "other_customers": [], "adjustments": [],
            "debt_total_latest_year": 3.0, "debt_current_portion_latest_year": 1.0,
            "covenant_max_net_debt_to_ebitda": 3.5,
            "narrative_excerpts": [{"title": "Outlook", "page": 2, "text": "We expect growth."}],
        })
        ext = _from_pdf(parsed, "acme.pdf")
        d = ext.data
        self.assertEqual(d.lines["vendas"], {2024: 10000.0, 2025: 12500.0})
        self.assertEqual(d.lines["cmv"], {2024: -6000.0, 2025: -7000.0})
        self.assertEqual(d.lines["map_tot"], {2025: 3000.0})
        self.assertEqual(d.lines["cov_lim"], {2025: 3.5})        # a ratio is not scaled
        self.assertEqual(d.sources["vendas"], "acme.pdf p.3")
        self.assertIn("We expect growth.", ext.source_text)


if __name__ == "__main__":
    unittest.main()
