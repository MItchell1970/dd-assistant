"""The review done in the app (status and comment) must reach the CSV and the Excel."""
import csv
import io
import unittest
from pathlib import Path

from openpyxl import load_workbook

from dd_assistant.input_loader import load_input
from dd_assistant.metrics import compute_metrics
from dd_assistant.report import COLUMNS, csv_bytes, write_excel
from dd_assistant.rules import alerts, evaluate

TEMPLATE = Path(__file__).resolve().parents[1] / "data" / "DD_template_modulo1.xlsx"
REVIEWS = {"R01": ("Accepted", "Customer A contract requested")}


class ReportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = load_input(TEMPLATE)
        cls.metrics = compute_metrics(cls.data)
        cls.results = evaluate(cls.data, cls.metrics)
        cls.first = alerts(cls.results)[0].rule.id     # R01 for Duriva

    def test_csv_keeps_review(self):
        raw = csv_bytes(self.results, REVIEWS)
        self.assertTrue(raw.startswith("﻿".encode("utf-8")))     # BOM: Excel reads accented characters
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(rows[0], COLUMNS)
        self.assertEqual(rows[1][1], self.first)
        self.assertEqual(rows[1][-2:], ["Accepted", "Customer A contract requested"])
        self.assertEqual(rows[2][-2:], ["Pending", ""])                 # not reviewed = Pending

    def test_excel_keeps_review(self):
        buf = io.BytesIO()
        write_excel(self.results, self.metrics, self.data.years, buf, REVIEWS)
        ws = load_workbook(io.BytesIO(buf.getvalue()))["Red Flags"]
        self.assertEqual(ws["B2"].value, self.first)
        self.assertEqual((ws["K2"].value, ws["L2"].value), REVIEWS["R01"])
        self.assertEqual(ws["K3"].value, "Pending")

    def test_without_reviews_everything_pending(self):
        raw = csv_bytes(self.results).decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(raw), delimiter=";"))[1:]
        self.assertEqual({r[-2] for r in rows}, {"Pending"})


if __name__ == "__main__":
    unittest.main()
